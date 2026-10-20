"""Prestação montada a partir das peças no Benner e dos resultados no tribunal."""

from __future__ import annotations

import io
from datetime import UTC, date, datetime

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.models.faturamento import BasaAtoCobrado, BasaFicha, MotorExecucao, MotorOportunidade
from app.models.tenant import Tenant
from app.modules.m12_faturamento.conferencia import DocumentoBenner
from app.modules.m12_faturamento.prestacao_auto import guardar_documentos, montar

PA = "08007677620248140073"
AM = "00000748420168047801"
FBS = "Fonseca Brasil Advogados"


def _d(cnj: str, dia: int, nome: str, mes: int = 9) -> DocumentoBenner:
    return DocumentoBenner(
        data=date(2026, mes, dia),
        nome=nome,
        tipo="Petição",
        ficha="CIVE.1",
        cnj=cnj,
        incluido_por=FBS,
        arquivo=f"{nome}.pdf",
    )


def test_rascunho(db: Session, as_tenant) -> None:  # type: ignore[no-untyped-def]
    t = Tenant(name="FBS rascunho", kind="law_firm")
    db.add(t)
    db.commit()
    with as_tenant(t.id):
        db.add_all(
            [
                BasaFicha(
                    tenant_id=t.id,
                    ficha="CIVE.1",
                    cnj=PA,
                    situacao="Em andamento",
                    adverso="Fulano",
                    importado_em=datetime.now(UTC),
                ),
                BasaAtoCobrado(
                    tenant_id=t.id,
                    competencia="2026-08",
                    cnj=AM,
                    ficha="CIVE.2",
                    ato_id="peticao_interlocutoria",
                    ato_texto="Petições interlocutórias",
                    data=date(2026, 8, 30),
                    comarca="Urucará",
                    acao="Execução",
                ),
            ]
        )
        ex = MotorExecucao(
            tenant_id=t.id, iniciado_em=datetime.now(UTC), referencia=date(2026, 10, 1), status="ok"
        )
        db.add(ex)
        db.commit()
        db.add(
            MotorOportunidade(
                tenant_id=t.id,
                execucao_id=ex.id,
                cnj=PA,
                evento="mandado_citacao_positivo",
                data_evento=date(2026, 9, 12),
                ato_id="citacao_efetivada",
                valor=130,
                papel="detecta",
                confianca="alta",
                status="ABERTA",
                movimento="Mandado cumprido",
            )
        )
        db.commit()
        g = guardar_documentos(
            db,
            t.id,
            [
                _d(PA, 10, "Petição interlocutória"),
                _d(PA, 10, "Petição interlocutória"),  # duplicado no export: guarda 1
                _d(AM, 5, "Petição interlocutória"),  # já cobrada em 30/08: fica fora
                _d(PA, 20, "Contrarrazões ao recurso de apelação"),
                _d(PA, 25, "Petição interlocutória", mes=8),  # fora do mês
            ],
        )
        assert g == {"novos": 4, "total": 4}
        x, rs = montar(db, "2026-09")
        assert rs["linhas"] == 3 and rs["a_conferir"] == 1  # a citação pede confirmação
        rows = list(load_workbook(io.BytesIO(x)).active.iter_rows(values_only=True))
        assert rows[0][:9] == (
            "NPC",
            "Nº Processo",
            "Comarca",
            "UF",
            "Ação",
            "Parte Adversa",
            "Ato processual realizado",
            "Data do protocolo",
            "Valor",
        )
        atos = sorted(r[6] for r in rows[1:])
        assert len(atos) == 3 and any("Contrarraz" in a for a in atos)
        assert all(r[3] == "Pará" and r[5] == "Fulano" for r in rows[1:])
