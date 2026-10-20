"""M12 — conferente do Benner: ato cobrado × documento no Benner, e peças sem cobrança."""

from __future__ import annotations

import io
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

from app.modules.m12_faturamento.auditoria import LinhaAto, auditar
from app.modules.m12_faturamento.conferencia import (
    DocumentoBenner,
    classificar,
    conferir,
    ler_documentos,
)
from app.modules.m12_faturamento.planilhas import gerar_excel

CNJ_A = "0000066-32.2015.8.04.4701"
CNJ_B = "0000079-31.2015.8.04.4701"
FBS = "Fonseca Brasil Advogados"


def _l(**kw: object) -> LinhaAto:
    base: dict[str, object] = dict(
        ficha="CIVE.1/2015",
        cnj=CNJ_A,
        comarca="Itacoatiara",
        uf="Amazonas",
        acao="Execução",
        parte_adversa="X",
        ato_texto="Petições interlocutórias",
        data=date(2026, 8, 10),
        valor=120.0,
    )
    base.update(kw)
    return LinhaAto(**base)  # type: ignore[arg-type]


def _d(
    nome: str,
    arquivo: str = "",
    dia: int = 10,
    cnj: str = CNJ_A,
    por: str = FBS,
    tipo: str = "Petição",
) -> DocumentoBenner:
    return DocumentoBenner(
        data=date(2026, 8, dia),
        nome=nome,
        tipo=tipo,
        ficha="CIVE.1/2015",
        cnj=cnj.replace(".", "").replace("-", ""),
        incluido_por=por,
        arquivo=arquivo,
    )


def _conf(linhas: list[LinhaAto], docs: list[DocumentoBenner], **kw: object):  # type: ignore[no-untyped-def]
    return conferir(auditar(linhas, "2026-08"), docs, "2026-08", **kw)  # type: ignore[arg-type]


def test_classifica_pelo_arquivo_com_rotulo_do_contrato() -> None:
    assert classificar(
        "manifestação", "Manifestação", "31.08.2026. Petições interlocutórias. Proc nº 1.pdf"
    ) == ("peca", "peticao_interlocutoria")
    assert classificar("COMPROVANTE DE PROTOCOLO", "Comprovante", "x.pdf") == ("protocolo", None)
    assert (
        classificar("Comprovante de custas", "Comprovante", "COMPROVANTE CARTA PRECATÓRIA.pdf")[0]
        == "anexo"
    )
    assert classificar("Manifestação", "Manifestação", "123.pdf") == ("generico", None)


def test_ok_generico_e_sem_documento() -> None:
    r = _conf(
        [
            _l(),  # interlocutória com a peça
            _l(cnj=CNJ_B, ficha="CIVE.2/2015"),  # só um "comprovante de protocolo"
            _l(ato_texto="Recursos apelação", data=date(2026, 8, 20), valor=250.0),  # nada
        ],
        [
            _d("Petição interlocutória", "10.08.2026 - Petição interlocutória.pdf"),
            _d("COMPROVANTE DE PROTOCOLO", "x.pdf", cnj=CNJ_B),
        ],
    )
    st = [c.status for c in r.atos]
    assert st == ["DOC_OK", "DOC_GENERICO", "SEM_DOCUMENTO"]
    # a pendência vira alerta na própria pré-auditoria
    assert any(a.codigo == "SEM_DOCUMENTO" for a in r.atos[2].auditada.achados)
    assert r.resumo()["por_status"]["SEM_DOCUMENTO"]["valor"] == 250.0


def test_um_documento_sustenta_um_ato_so() -> None:
    r = _conf(
        [_l(), _l(data=date(2026, 8, 12))],
        [_d("Petição interlocutória", "Petição interlocutória.pdf")],
    )
    assert sorted(c.status for c in r.atos) == ["DOC_OK", "SEM_DOCUMENTO"]
    assert "já sustentam outros atos" in next(c for c in r.atos if c.documento is None).detalhe


def test_documento_de_terceiro_so_vale_se_for_a_propria_peca() -> None:
    r = _conf(
        [
            _l(ato_texto="Ajuizamento", acao="Execução de título extrajudicial", valor=200.0),
            _l(cnj=CNJ_B, ficha="CIVE.2/2015"),
        ],
        [
            _d("Petição inicial", "inicial.pdf", por="PREPOSTO DO BANCO"),
            _d("Despacho", "despacho.pdf", cnj=CNJ_B, por="PREPOSTO DO BANCO", tipo="Despacho"),
        ],
    )
    assert r.atos[0].status == "DOC_OK" and r.atos[0].terceiro
    assert r.atos[1].status == "SEM_DOCUMENTO" and r.atos[1].terceiros == 1


def test_peca_na_ficha_errada_casa_pelo_cnj_citado() -> None:
    r = _conf(
        [_l(cnj=CNJ_B, ficha="CIVE.2/2015")],
        [_d("Petição interlocutória", f"10.08.2026 Petição interlocutória {CNJ_B}.pdf", cnj=CNJ_A)],
    )
    assert r.atos[0].status == "DOC_OK"
    assert "cita o processo" in r.atos[0].detalhe


def test_nao_cobrados_respeitam_outras_competencias_e_atos_de_resultado() -> None:
    docs = [
        _d("Contrarrazões", "05.08.2026. Contrarrazões em recursos apelação.pdf", dia=5),
        _d("Petição interlocutória", "Petição interlocutória.pdf", cnj=CNJ_B, dia=6),
        _d("Manifestação requerendo a extinção", "pedido de extinção.pdf", dia=7),
        _d("Petição interlocutória", "Petição interlocutória.pdf", dia=10),
    ]
    r = _conf([_l()], docs, cobrados_outras=[(CNJ_B, "peticao_interlocutoria", date(2026, 7, 30))])
    nc = {(n.ato_id, n.documento.cnj[-4:]) for n in r.nao_cobrados}
    # contrarrazões ficou sem cobrança; a interlocutória do B já foi cobrada em julho;
    # pedido de extinção não é o ato de resultado; a interlocutória do dia 10 sustenta a linha
    assert nc == {("contrarrazoes", "4701")}
    assert r.nao_cobrados[0].valor == 200 and r.nao_cobrados[0].coluna == "rec_cred"


def test_ler_export_e_gerar_abas(tmp_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.append(
        [
            "Data do documento",
            "Nome",
            "Tipo documento",
            "Ficha",
            "Nº Processo principal",
            "Incluído por",
            "Arquivo",
        ]
    )
    ws.append(
        [
            datetime(2026, 8, 10),
            "Petição interlocutória",
            "Petição",
            "CIVE.1/2015",
            CNJ_A,
            FBS,
            "p.pdf",
        ]
    )
    ws.append(
        [
            datetime(2026, 8, 10),
            "Petição interlocutória",
            "Petição",
            "CIVE.1/2015",
            CNJ_A,
            FBS,
            "p.pdf",
        ]
    )
    buf = io.BytesIO()
    wb.save(buf)
    docs = ler_documentos([buf.getvalue(), buf.getvalue()])
    assert len(docs) == 1  # duplicados entre exports caem
    aud = auditar([_l()], "2026-08")
    res = conferir(aud, docs, "2026-08")
    out = tmp_path / "c.xlsx"
    rs = gerar_excel(aud, "2026-08", out, conferencia=res)
    assert rs["conferencia"]["por_status"]["DOC_OK"]["atos"] == 1
    wb2 = load_workbook(out)
    assert wb2.sheetnames[:3] == ["Resumo", "Conferência Benner", "Não cobrados"]


def test_endpoint_conferencia(engine, monkeypatch) -> None:  # type: ignore[no-untyped-def]
    from fastapi.testclient import TestClient

    from app.core import db as core_db
    from app.core.config import get_settings
    from app.main import app

    core_db.SessionLocal.configure(bind=engine)
    monkeypatch.setattr(get_settings(), "ops_key", "k")
    c = TestClient(app)
    tid = c.post("/faturamento/basa/setup", headers={"X-Ops-Key": "k"}).json()["tenant_id"]

    def xlsx(rows: list[list[object]]) -> bytes:
        wb = Workbook()
        ws = wb.active
        assert ws is not None
        for r in rows:
            ws.append(r)
        b = io.BytesIO()
        wb.save(b)
        return b.getvalue()

    prest = xlsx(
        [
            [
                "NPC",
                "Nº Processo",
                "Comarca",
                "UF",
                "Ação",
                "Parte Adversa",
                "Ato processual realizado",
                "Data do protocolo",
                "Valor",
            ],
            [
                "CIVE.1/2015",
                CNJ_A,
                "Itacoatiara",
                "Amazonas",
                "Execução",
                "X",
                "Petições interlocutórias",
                "2026-08-10",
                120,
            ],
        ]
    )
    docs = xlsx(
        [
            [
                "Data do documento",
                "Nome",
                "Tipo documento",
                "Ficha",
                "Nº Processo principal",
                "Incluído por",
                "Arquivo",
            ],
            [
                datetime(2026, 8, 10),
                "Petição interlocutória",
                "Petição",
                "CIVE.1/2015",
                CNJ_A,
                FBS,
                "p.pdf",
            ],
        ]
    )
    h = {"X-Ops-Key": "k", "X-Tenant-ID": tid}
    files = [("prestacao", ("p.xlsx", prest)), ("documentos", ("d.xlsx", docs))]
    r = c.post(
        "/faturamento/basa/conferencia?competencia=2026-08&formato=json", headers=h, files=files
    )
    assert r.status_code == 200, r.text
    assert r.json()["conferencia"]["por_status"] == {"DOC_OK": {"atos": 1, "valor": 120.0}}
    r = c.post("/faturamento/basa/conferencia?competencia=2026-08", headers=h, files=files)
    assert r.status_code == 200 and r.content[:2] == b"PK"
