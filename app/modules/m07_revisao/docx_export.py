"""Exportação da minuta aprovada para DOCX (python-docx) com bloco de capa obrigatório.

Só é chamado depois de `assert_exportable()`; o bloco de capa identifica a origem (playbook,
versão, OAB que aprovou as teses) e a revisão (OAB do revisor, tempo de revisão).
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass
from datetime import datetime

_BULLET = re.compile(r"^\s*[-*]\s+(.*)$")
_NUMBERED = re.compile(r"^\s*\d+[.)]\s+(.*)$")
_HEADING = re.compile(r"^(#{1,6})\s+(.*)$")
_QUOTE = re.compile(r"^>\s?(.*)$")
_INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|«[^»]+»)")


@dataclass
class CoverInfo:
    generated_at: datetime
    playbook_id: str | None
    playbook_version: int | None
    playbook_approved_by_oab: str | None
    reviewer_oab: str
    approved_at: datetime
    review_seconds: int | None
    cnj_number: str
    draft_type: str
    draft_version: int

    def cover_line(self) -> str:
        pb = (
            f"playbook {self.playbook_id} v{self.playbook_version} aprovado por OAB/"
            f"{_fmt_oab(self.playbook_approved_by_oab)}"
            if self.playbook_id
            else "sem playbook vinculado"
        )
        ts = self.generated_at.strftime("%Y-%m-%d %H:%M UTC")
        return f"MINUTA — revisão obrigatória — gerada em {ts} — {pb}"


def _fmt_oab(oab: str | None) -> str:
    """'PA 12724' → 'PA 12.724'."""
    if not oab:
        return "—"
    m = re.match(r"^([A-Z]{2})\s?(\d+)([A-Z]?)$", oab.strip().upper())
    if not m:
        return oab
    uf, num, suf = m.groups()
    return f"{uf} {int(num):,}".replace(",", ".") + suf


def _add_inline(paragraph, text: str) -> None:  # type: ignore[no-untyped-def]
    pos = 0
    for m in _INLINE.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos : m.start()])
        tok = m.group(0)
        if tok.startswith("**"):
            paragraph.add_run(tok[2:-2]).bold = True
        elif tok.startswith("*"):
            paragraph.add_run(tok[1:-1]).italic = True
        else:  # «lacuna para o revisor»
            run = paragraph.add_run(tok)
            run.bold = True
            run.font.highlight_color = 7  # WD_COLOR_INDEX.YELLOW
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def render_docx(body_markdown: str, cover: CoverInfo) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    # --- bloco de capa -------------------------------------------------------
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(cover.cover_line())
    r.bold = True
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run(
        f"Processo {cover.cnj_number} · {cover.draft_type} v{cover.draft_version} · "
        f"aprovada por OAB/{_fmt_oab(cover.reviewer_oab)} em "
        f"{cover.approved_at.strftime('%Y-%m-%d %H:%M UTC')}"
        + (f" após {cover.review_seconds} s de revisão" if cover.review_seconds else "")
    ).italic = True
    doc.add_paragraph(
        "Documento gerado por sistema como sugestão; a estratégia e o conteúdo são de "
        "responsabilidade do advogado que revisou e aprovou. Citações [Pn] remetem às páginas "
        "dos autos/subsídios registradas em draft_citations."
    ).runs[0].font.size = Pt(9)
    doc.add_paragraph()

    # --- corpo (Markdown simples) ---------------------------------------------
    in_quote = False
    for raw in body_markdown.splitlines():
        line = raw.rstrip()
        if not line.strip():
            in_quote = False
            continue
        if m := _HEADING.match(line):
            level = min(len(m.group(1)), 4)
            doc.add_heading(m.group(2).strip(), level=level)
            continue
        if m := _QUOTE.match(line):
            q = doc.add_paragraph(style="Intense Quote" if not in_quote else "Quote")
            _add_inline(q, m.group(1).strip())
            in_quote = True
            continue
        if m := _BULLET.match(line):
            _add_inline(doc.add_paragraph(style="List Bullet"), m.group(1))
            continue
        if m := _NUMBERED.match(line):
            _add_inline(doc.add_paragraph(style="List Number"), m.group(1))
            continue
        para = doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        _add_inline(para, line.strip())

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


__all__ = ["CoverInfo", "render_docx"]
