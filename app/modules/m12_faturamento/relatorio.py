"""Seleção das oportunidades (detecção + conciliação) e o relatório .xlsx do motor."""

from __future__ import annotations

import io
from collections import Counter, defaultdict
from collections.abc import Iterable
from datetime import date, timedelta
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from app.modules.m12_faturamento.deteccao import (
    DIAS_PARA_ATENDER,
    Oportunidade,
    Processo,
    conciliar,
    detectar,
)
from app.modules.m12_faturamento.tabela import carregar_tabela


def selecionar(
    processos: Iterable[Processo],
    cobrados: dict[str, list[tuple[date, str]]],
    hoje: date,
    cobertura: tuple[date, date],
) -> list[Oportunidade]:
    """Detecta e concilia. `cobertura` = período das prestações de contas conhecidas: eventos
    cuja janela de atendimento passa do fim dela continuam na fila (ABERTA), nunca PERDIDA."""
    ini, fim = cobertura
    todas: list[Oportunidade] = []
    for proc in processos:
        ops = detectar(proc, hoje, desde=ini - timedelta(days=DIAS_PARA_ATENDER))
        conciliar(ops, cobrados, hoje)
        for o in ops:
            if o.status == "ALERTA" or o.evento == "processo_sem_evento_90_dias":
                pass
            elif o.data_evento < ini - timedelta(days=3) and o.status != "ATENDIDA":
                continue  # evento anterior à cobertura e não atendido: sem como julgar
            elif (
                o.status == "PERDIDA"
                and (o.prazo or o.data_evento + timedelta(days=DIAS_PARA_ATENDER)) > fim
            ):
                o.status = "ABERTA"
            todas.append(o)
    return todas


def gerar_relatorio(
    ops: list[Oportunidade],
    n_processos: int,
    fichas: dict[str, str],
    hoje: date,
    cobertura: tuple[date, date],
) -> tuple[bytes, dict[str, Any]]:
    """Planilha do motor (bytes) + resumo (JSON) a partir das oportunidades selecionadas."""
    tab = carregar_tabela()
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Resumo"
    bold = Font(bold=True)
    ws.append(
        [
            f"Motor de oportunidades BASA — {hoje:%d/%m/%Y} — cobertura das prestações "
            f"{cobertura[0]:%d/%m/%Y} a {cobertura[1]:%d/%m/%Y}"
        ]
    )
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([f"Processos lidos no DataJud: {n_processos}"])
    ws.append([])
    ws.append(
        [
            "Evento",
            "Ato",
            "Papel da IA",
            "Confiança",
            "Atendidas",
            "Perdidas",
            "Abertas (fila)",
            "R$ perdido",
            "R$ em aberto",
        ]
    )
    for c in ws[4]:
        c.font = bold
    agg: dict[tuple[str, str, str, str], Counter[str]] = defaultdict(Counter)
    val: dict[tuple[str, str, str, str], defaultdict[str, float]] = defaultdict(
        lambda: defaultdict(float)
    )
    for o in ops:
        if o.status == "ALERTA":
            continue
        k = (o.evento, o.ato_id or "", o.papel or "", o.confianca)
        agg[k][o.status] += 1
        val[k][o.status] += o.valor or 0
    linhas = sorted(agg, key=lambda k: -(val[k]["PERDIDA"] + val[k]["ABERTA"]))
    tot: defaultdict[str, float] = defaultdict(float)
    for k in linhas:
        nome_ato = tab.atos[k[1]].nome if k[1] in tab.atos else k[1]
        ws.append(
            [
                k[0],
                nome_ato,
                k[2],
                k[3],
                agg[k]["ATENDIDA"],
                agg[k]["PERDIDA"],
                agg[k]["ABERTA"],
                round(val[k]["PERDIDA"], 2),
                round(val[k]["ABERTA"], 2),
            ]
        )
        for s in ("ATENDIDA", "PERDIDA", "ABERTA"):
            tot[s] += agg[k][s]
            tot["v" + s] += val[k][s]
    ws.append(
        [
            "TOTAL",
            "",
            "",
            "",
            tot["ATENDIDA"],
            tot["PERDIDA"],
            tot["ABERTA"],
            round(tot["vPERDIDA"], 2),
            round(tot["vABERTA"], 2),
        ]
    )
    for c in ws[ws.max_row]:
        c.font = bold
    for i, w in enumerate([40, 46, 10, 10, 10, 10, 13, 13, 13], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    def aba(nome: str, filtro: Any) -> None:
        sh = wb.create_sheet(nome)
        cab = [
            "Processo",
            "Ficha",
            "Evento",
            "Data do evento",
            "Ato cabível",
            "Valor",
            "Papel da IA",
            "Prazo de referência",
            "Confiança",
            "Movimento",
            "Observação",
            "Atendida por",
        ]
        sh.append(cab)
        for c in sh[1]:
            c.font = bold
        for o in sorted(filter(filtro, ops), key=lambda x: (-(x.valor or 0), x.data_evento)):
            cnj = o.cnj
            fmt = f"{cnj[:7]}-{cnj[7:9]}.{cnj[9:13]}.{cnj[13]}.{cnj[14:16]}.{cnj[16:]}"
            sh.append(
                [
                    fmt,
                    fichas.get(cnj, ""),
                    o.evento,
                    o.data_evento,
                    tab.atos[o.ato_id].nome if o.ato_id in tab.atos else "",
                    o.valor,
                    o.papel,
                    o.prazo,
                    o.confianca,
                    o.movimento,
                    o.observacao,
                    o.atendida_por or "",
                ]
            )
        for i, w in enumerate([26, 18, 34, 12, 40, 9, 10, 12, 10, 40, 46, 26], start=1):
            sh.column_dimensions[get_column_letter(i)].width = w
        for row in sh.iter_rows(min_row=2):
            row[3].number_format = "DD/MM/YYYY"
            row[7].number_format = "DD/MM/YYYY"
        sh.freeze_panes = "A2"

    aba("Fila aberta", lambda o: o.status == "ABERTA" and o.evento != "processo_sem_evento_90_dias")
    aba("Perdidas", lambda o: o.status == "PERDIDA")
    aba("Sem movimento 90 dias", lambda o: o.evento == "processo_sem_evento_90_dias")
    aba("Alertas", lambda o: o.status == "ALERTA" and o.data_evento >= hoje - timedelta(days=120))
    aba("Atendidas", lambda o: o.status == "ATENDIDA")
    buf = io.BytesIO()
    wb.save(buf)
    alertas = Counter(o.evento for o in ops if o.status == "ALERTA")
    gat = [o for o in ops if o.status == "ABERTA" and o.evento != "processo_sem_evento_90_dias"]
    return buf.getvalue(), {
        # o número que importa: atos com gatilho real (sem os processos apenas parados)
        "abertas_com_gatilho": len(gat),
        "valor_aberto_com_gatilho": round(sum(o.valor or 0 for o in gat), 2),
        "parados_90_dias": sum(1 for o in ops if o.evento == "processo_sem_evento_90_dias"),
        "descartadas_pela_equipe": sum(1 for o in ops if o.status == "DESCARTADA"),
        "processos": n_processos,
        "atendidas": tot["ATENDIDA"],
        "perdidas": tot["PERDIDA"],
        "abertas": tot["ABERTA"],
        "valor_perdido": round(tot["vPERDIDA"], 2),
        "valor_aberto": round(tot["vABERTA"], 2),
        "alertas": dict(alertas),
        "por_evento": {
            k[0] + "|" + k[3]: {
                **agg[k],
                "R$perdido": round(val[k]["PERDIDA"]),
                "R$aberto": round(val[k]["ABERTA"]),
            }
            for k in linhas
        },
    }
