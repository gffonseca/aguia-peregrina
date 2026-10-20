"""Entrada e saída em Excel: atos praticados → pré-auditoria → Anexo II por estado.

Entrada aceita o layout da prestação já usada pelo escritório (NPC, Nº Processo, Comarca, UF,
Ação, Parte Adversa, Ato processual realizado, Data do protocolo, Valor) e ignora linhas de total.
Saída: um .xlsx com "Resumo", "Pré-auditoria" e uma aba "Anexo II - <UF>" por estado, no modelo
do contrato (Nº, Processo, Comarca, Matéria, Tipo de Ação, Parte Adversa, Ato, Data, Valor).
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.modules.m12_faturamento.auditoria import (
    FichaBenner,
    LinhaAto,
    LinhaAuditada,
    resumo,
    to_date,
)
from app.modules.m12_faturamento.tabela import ROTULO_COLUNA, normalizar

_CAMPOS = {
    "ficha": ("npc", "ficha", "pasta"),
    "cnj": ("n processo", "no processo", "processo", "numero do processo", "n processo principal"),
    "comarca": ("comarca",),
    "uf": ("uf", "estado"),
    "acao": ("acao", "tipo de acao"),
    "parte_adversa": ("parte adversa", "adverso", "adverso principal"),
    "ato_texto": ("ato processual realizado", "ato", "prazo ou ato processual realizado"),
    "data": ("data do protocolo", "data protocolo", "data"),
    "valor": ("valor", "valor do ato"),
}


def _mapear_cabecalho(cab: list[Any]) -> dict[str, int]:
    idx: dict[str, int] = {}
    norm = [normalizar(str(c or "")).replace("º", "").replace("°", "") for c in cab]
    for campo, nomes in _CAMPOS.items():
        for i, c in enumerate(norm):
            if c in nomes or any(c.startswith(n) for n in nomes):
                if i not in idx.values():
                    idx[campo] = i
                    break
    faltam = [c for c in ("ficha", "cnj", "ato_texto", "valor") if c not in idx]
    if faltam:
        raise ValueError(f"colunas obrigatórias não encontradas: {faltam}; cabeçalho: {cab}")
    return idx


def ler_atos(path: str | Path) -> list[LinhaAto]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    idx = _mapear_cabecalho(list(rows[0]))
    out: list[LinhaAto] = []

    def g(r: tuple[Any, ...], k: str) -> Any:
        i = idx.get(k)
        return r[i] if i is not None and i < len(r) else None

    for n, r in enumerate(rows[1:], start=2):
        ficha, cnj, ato = g(r, "ficha"), g(r, "cnj"), g(r, "ato_texto")
        if not (ficha or cnj or ato):
            continue  # linha de total ou vazia
        v = g(r, "valor")
        try:
            valor = float(v) if v not in (None, "") else None
        except (TypeError, ValueError):
            valor = None
        out.append(
            LinhaAto(
                ficha=str(ficha or "").strip(),
                cnj=str(cnj or "").strip(),
                comarca=str(g(r, "comarca") or "").strip(),
                uf=str(g(r, "uf") or "").strip(),
                acao=str(g(r, "acao") or "").strip(),
                parte_adversa=str(g(r, "parte_adversa") or "").strip(),
                ato_texto=str(ato or "").strip(),
                data=to_date(g(r, "data")),
                valor=valor,
                origem=n,
            )
        )
    return out


def ler_benner(paths: list[str | Path]) -> dict[str, FichaBenner]:
    """Lê os exports 'Fichas.xlsx' do Benner (Ficha, Nº processo principal, Situação)."""
    out: dict[str, FichaBenner] = {}
    for p in paths:
        ws = load_workbook(p, read_only=True, data_only=True).worksheets[0]
        rows = list(ws.iter_rows(values_only=True))
        cab = [normalizar(str(c or "")) for c in rows[0]]
        i_f = cab.index("ficha")
        i_c = next(i for i, c in enumerate(cab) if c.startswith("n") and "processo" in c)
        i_s = cab.index("situacao")
        for r in rows[1:]:
            if r[i_f]:
                out[str(r[i_f]).strip()] = FichaBenner(
                    ficha=str(r[i_f]).strip(), cnj=str(r[i_c] or ""), situacao=str(r[i_s] or "")
                )
    return out


_FILL = {
    "OK": PatternFill("solid", fgColor="E5F4E4"),
    "ALERTA": PatternFill("solid", fgColor="F9EEDA"),
    "BLOQUEIO": PatternFill("solid", fgColor="FBEBEB"),
}
_BOLD = Font(bold=True)


def _cabecalho(ws: Any, cols: list[str], larguras: list[int]) -> None:
    ws.append(cols)
    for i, w in enumerate(larguras, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
        c = ws.cell(row=1, column=i)
        c.font = _BOLD
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"


def _materia(a: LinhaAuditada) -> str:
    return "Trabalhista" if a.coluna == "trabalhista" else "Cível"


def gerar_excel(
    auditadas: list[LinhaAuditada], competencia: str, destino: str | Path
) -> dict[str, Any]:
    wb = Workbook()
    rs = resumo(auditadas)

    ws = wb.active
    assert ws is not None
    ws.title = "Resumo"
    ws.append([f"Prestação de contas BASA FB1798 — competência {competencia}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([])
    ws.append(
        [
            "UF",
            "Linhas",
            "OK",
            "Alerta",
            "Bloqueio",
            "Declarado (R$)",
            "Faturável no Anexo II (R$)",
            "Retido por bloqueio (R$)",
        ]
    )
    for c in ws[3]:
        c.font = _BOLD
    for uf, u in rs["por_uf"].items():
        ws.append(
            [
                uf,
                int(u.get("linhas", 0)),
                int(u.get("n_OK", 0)),
                int(u.get("n_ALERTA", 0)),
                int(u.get("n_BLOQUEIO", 0)),
                round(u.get("declarado", 0), 2),
                round(u.get("faturavel", 0), 2),
                round(u.get("bloqueado", 0), 2),
            ]
        )
    ws.append([])
    ws.append(["Regra", "Ocorrências"])
    ws.cell(row=ws.max_row, column=1).font = _BOLD
    ws.cell(row=ws.max_row, column=2).font = _BOLD
    for regra, n in rs["regras"].items():
        ws.append([regra, n])
    for i, w in enumerate([28, 10, 8, 8, 10, 16, 22, 20], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    pa = wb.create_sheet("Pré-auditoria")
    _cabecalho(
        pa,
        [
            "Linha",
            "Status",
            "Achados",
            "Ficha",
            "Processo",
            "UF",
            "Ação",
            "Ato informado",
            "Ato na tabela",
            "Coluna da tabela",
            "Data",
            "Valor declarado",
            "Valor tabela",
            "Ato de resultado",
            "Ficha correta (Benner)",
        ],
        [7, 10, 60, 18, 26, 5, 26, 40, 40, 20, 11, 12, 12, 10, 18],
    )
    for a in sorted(
        auditadas, key=lambda x: (-["OK", "ALERTA", "BLOQUEIO"].index(x.status), x.linha.origem)
    ):
        pa.append(
            [
                a.linha.origem,
                a.status,
                "; ".join(f"{x.codigo}: {x.detalhe}" for x in a.achados),
                a.linha.ficha,
                a.linha.cnj,
                a.linha.uf,
                a.linha.acao,
                a.linha.ato_texto,
                a.ato_nome or "",
                ROTULO_COLUNA.get(a.coluna, a.coluna),
                a.linha.data,
                a.linha.valor,
                a.valor_tabela,
                "sim" if a.resultado else "",
                a.ficha_sugerida or "",
            ]
        )
        pa.cell(row=pa.max_row, column=2).fill = _FILL[a.status]
        pa.cell(row=pa.max_row, column=11).number_format = "DD/MM/YYYY"

    for uf in sorted({a.linha.uf for a in auditadas}):
        ax = wb.create_sheet(f"Anexo II - {uf}")
        _cabecalho(
            ax,
            [
                "Nº",
                "Processo",
                "Comarca",
                "Matéria",
                "Tipo de Ação",
                "Parte Adversa",
                "Prazo ou Ato Processual Realizado",
                "Data Protocolo",
                "Valor do Ato",
                "Ficha",
            ],
            [6, 26, 18, 12, 26, 34, 50, 13, 12, 18],
        )
        linhas = sorted(
            (a for a in auditadas if a.linha.uf == uf and a.status != "BLOQUEIO"),
            key=lambda x: (x.linha.data or competencia, x.linha.cnj),
        )
        total = 0.0
        for n, a in enumerate(linhas, start=1):
            v = a.valor_faturavel
            total += v
            ax.append(
                [
                    n,
                    a.linha.cnj,
                    a.linha.comarca,
                    _materia(a),
                    a.linha.acao,
                    a.linha.parte_adversa,
                    a.ato_nome,
                    a.linha.data,
                    v,
                    a.ficha_sugerida or a.linha.ficha,
                ]
            )
            ax.cell(row=ax.max_row, column=8).number_format = "DD/MM/YYYY"
            ax.cell(row=ax.max_row, column=9).number_format = "#,##0.00"
        ax.append([None, None, None, None, None, None, "TOTAL", None, round(total, 2)])
        ax.cell(row=ax.max_row, column=7).font = _BOLD
        ax.cell(row=ax.max_row, column=9).font = _BOLD
        ax.cell(row=ax.max_row, column=9).number_format = "#,##0.00"

    Path(destino).parent.mkdir(parents=True, exist_ok=True)
    wb.save(destino)
    return rs
