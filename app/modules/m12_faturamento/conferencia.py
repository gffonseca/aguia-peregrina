"""Conferente do Benner: cada ato cobrado tem a peça no Benner? E o que subiu sem ser cobrado?

Entrada: a prestação já pré-auditada (`LinhaAuditada`) e o export "Pesquisa de documentos" do
Benner (Data do documento, Nome, Tipo documento, Ficha, Nº Processo principal, Incluído por,
Arquivo). Só contam como prova os documentos incluídos pelo escritório (`proprios`); os
demais (prepostos do Banco, outros escritórios, carga do sistema) aparecem como "de terceiro".

Para cada ato cobrado, procura documento do mesmo processo numa janela em torno da data do
protocolo e classifica:

- DOC_OK: peça do mesmo tipo do ato (ex.: apelação ↔ "Apelação.pdf").
- DOC_GENERICO: há peça/protocolo do escritório, mas o nome não diz o tipo ("Manifestação",
  "Comprovante de protocolo") — basta o revisor abrir.
- DOC_DIVERGENTE: só há documento de outro tipo (ex.: cobrou contestação, subiu custas).
- DOC_FORA_DA_JANELA: há documento livre do escritório no processo, mas longe da data cobrada.
- SEM_DOCUMENTO: nenhuma peça livre do escritório no processo (ou todas já sustentam outros
  atos) — subir a peça antes de enviar a prestação.

Um documento sustenta um único ato (dois protocolos cobrados pedem duas peças). Passadas em
ordem de força: primeiro casamentos exatos, depois genéricos, depois divergentes.

Na direção inversa, peças do escritório cujo nome reconhece um ato remunerado da tabela, no mês
da competência, sem ato cobrado correspondente (mesmo processo e ato em ±45 dias em qualquer
prestação carregada) viram "possível ato não cobrado" — receita que ficou na mesa.
"""

from __future__ import annotations

import re
from collections import defaultdict
from collections.abc import Iterable, Sequence
from dataclasses import dataclass, field
from datetime import date, timedelta
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.core import cnj as cnjlib
from app.modules.m12_faturamento.auditoria import Achado, LinhaAuditada, to_date
from app.modules.m12_faturamento.tabela import (
    carregar_tabela,
    enquadrar_coluna,
    nao_remunerado,
    normalizar,
    reconhecer_ato,
)

PROPRIOS_PADRAO = ("fonseca brasil",)
ANTES_DIAS = 10  # peça datada até 10 dias antes do protocolo (assinada antes de protocolar)
DEPOIS_DIAS = 20  # ou até 20 dias depois (upload tardio)
JANELA_REPETICAO = 45  # mesmo ato no mesmo processo já cobrado nesse intervalo

_PROTOCOLO = re.compile(r"protocolo")
_ANEXO = re.compile(
    r"comprovante|custas|preparo|guia|formulario|analise de risco|calculo|extrato|"
    r"planilha|procuracao|substabelecimento|contrato|lote de documentos|dispensa recursal"
)
_SENTENCA = re.compile(r"sentenca|acordao|homologa")
_GENERICO = re.compile(r"^(manifestacao|peticao|peticoes|juntada|lote de documentos)$")
# atos próximos: a peça de um costuma vir nomeada como o outro
_FAMILIAS = (
    {"peticao_interlocutoria", "juntada_planilha_custas", "diligencia_cartorio"},
    {"apelacao", "contrarrazoes"},
    {"impugnacao_embargos_execucao", "impugnacao_cumprimento", "impugnacao_calculos"},
    {"audiencia_conciliacao", "audiencia_instrucao", "audiencia_jec"},
)
STATUS_ORDEM = (
    "SEM_DOCUMENTO",
    "DOC_DIVERGENTE",
    "DOC_FORA_DA_JANELA",
    "DOC_GENERICO",
    "DOC_OK",
)


def _dig(s: Any) -> str:
    return re.sub(r"\D", "", str(s or ""))


def cnj_fmt(d: str | None) -> str:
    d = _dig(d)
    if len(d) != 20:
        return d
    return f"{d[:7]}-{d[7:9]}.{d[9:13]}.{d[13]}.{d[14:16]}.{d[16:]}"


@dataclass
class DocumentoBenner:
    data: date | None
    nome: str
    tipo: str
    ficha: str
    cnj: str  # só dígitos
    incluido_por: str
    arquivo: str
    origem: int = 0
    classe: str = ""  # peca | generico | protocolo | anexo
    ato_id: str | None = None
    cnj_citado: str | None = None  # CNJ escrito no nome/arquivo, quando difere do da ficha

    def __post_init__(self) -> None:
        if not self.classe:
            self.classe, self.ato_id = classificar(self.nome, self.tipo, self.arquivo)
        c = cnj_no_texto(self.arquivo, self.nome)
        if c and c != self.cnj and (pc := cnjlib.parse(c)) is not None and pc.is_valid:
            self.cnj_citado = c

    @property
    def processos(self) -> tuple[str, ...]:
        """A peça pode estar na ficha do principal citando o recurso/incidente (ou o contrário)."""
        return tuple(x for x in (self.cnj, self.cnj_citado) if x)

    def proprio(self, proprios: Sequence[str]) -> bool:
        q = normalizar(self.incluido_por)
        return any(p in q for p in proprios)

    @property
    def rotulo(self) -> str:
        d = self.data.strftime("%d/%m/%Y") if self.data else "s/ data"
        return f"{d} · {self.nome or self.tipo} [{self.arquivo}]"


_CNJ_TXT = re.compile(
    r"(\d{7})\s*-?\s*(\d{2})\s*\.?\s*(\d{4})\s*\.?\s*(\d)\s*\.?\s*(\d{2})\s*\.?\s*(\d{4})"
)
_ANEXO_INICIO = re.compile(r"^(comprovante(?! de protocolo)|custas|guia|preparo|boleto|extrato)")


def cnj_no_texto(*textos: str) -> str | None:
    """Primeiro número CNJ citado no nome/arquivo do documento (só dígitos)."""
    for t in textos:
        m = _CNJ_TXT.search(t or "")
        if m:
            return "".join(m.groups())
    return None


def classificar(nome: str, tipo: str, arquivo: str) -> tuple[str, str | None]:
    """Documento → (classe, ato da tabela).

    O escritório nomeia o arquivo com o rótulo do ato do contrato ("31.08.2026. Petições
    interlocutórias. Proc nº ..."), então o arquivo vem primeiro; depois o nome; o tipo
    cadastrado no Benner é o menos confiável.
    """
    n = normalizar(nome)
    a = normalizar(re.sub(r"\.\w{2,4}$", "", arquivo or ""))
    t = normalizar(tipo)
    ato_a = None if _ANEXO_INICIO.match(a) else reconhecer_ato(a)
    ato_n = None if _ANEXO_INICIO.match(n) else reconhecer_ato(n)
    ato = ato_a or ato_n
    if ato == "diligencia_cartorio" and (_ANEXO.search(n) or _ANEXO.search(a)):
        ato = None
    if ato:
        return "peca", ato
    if _PROTOCOLO.search(n) or _PROTOCOLO.search(a):
        return "protocolo", None
    if _ANEXO.search(n) or _ANEXO.search(t) or _ANEXO_INICIO.match(a):
        return "anexo", None
    ato_t = reconhecer_ato(t)
    if ato_t and ato_t != "diligencia_cartorio":
        return "peca", ato_t
    return "generico", None


def ler_documentos(paths: Iterable[str | Path | bytes]) -> list[DocumentoBenner]:
    """Lê um ou mais exports "Pesquisa de documentos" (aceita caminho ou bytes)."""
    import io

    out: list[DocumentoBenner] = []
    for p in paths:
        src = io.BytesIO(p) if isinstance(p, bytes) else p
        ws = load_workbook(src, read_only=True, data_only=True).active
        assert ws is not None
        it = ws.iter_rows(values_only=True)
        cab = [normalizar(str(c or "")) for c in next(it)]

        def col(*nomes: str, _cab: list[str] = cab) -> int:
            for nm in nomes:
                if nm in _cab:
                    return _cab.index(nm)
            raise ValueError(f"coluna ausente no export de documentos: {nomes[0]}")

        i_data = col("data do documento", "data")
        i_nome = col("nome")
        i_tipo = col("tipo documento", "tipo")
        i_ficha = col("ficha")
        i_cnj = col("no processo principal", "n processo principal", "processo")
        i_inc = col("incluido por")
        i_arq = col("arquivo")
        for n, r in enumerate(it, start=2):
            if not any(r):
                continue
            out.append(
                DocumentoBenner(
                    data=to_date(r[i_data]),
                    nome=str(r[i_nome] or "").strip(),
                    tipo=str(r[i_tipo] or "").strip(),
                    ficha=str(r[i_ficha] or "").strip(),
                    cnj=_dig(r[i_cnj]),
                    incluido_por=str(r[i_inc] or "").strip(),
                    arquivo=str(r[i_arq] or "").strip(),
                    origem=n,
                )
            )
    # o mesmo documento pode vir em dois exports (meses sobrepostos)
    vistos: set[tuple[Any, ...]] = set()
    unicos = []
    for d in out:
        k = (d.cnj, d.ficha, d.data, d.nome, d.arquivo, d.incluido_por)
        if k not in vistos:
            vistos.add(k)
            unicos.append(d)
    return unicos


@dataclass
class ConferenciaAto:
    auditada: LinhaAuditada
    status: str
    documento: DocumentoBenner | None = None
    outros: list[DocumentoBenner] = field(default_factory=list)  # demais do processo, p/ revisor
    terceiros: int = 0
    terceiro: bool = False  # o documento que sustenta foi incluído por outro usuário
    ja_usados: int = 0  # peças do processo que já sustentam outros atos cobrados

    @property
    def detalhe(self) -> str:
        if self.documento:
            quem = f" — incluído por {self.documento.incluido_por}" if self.terceiro else ""
            outro = (
                f" — cita o processo {cnj_fmt(self.documento.cnj_citado)}"
                if self.documento.cnj_citado
                else ""
            )
            return self.documento.rotulo + quem + outro
        if self.outros:
            return "mais próximo: " + self.outros[0].rotulo
        if self.ja_usados:
            return (
                f"as {self.ja_usados} peça(s) do escritório no processo já sustentam outros atos "
                "cobrados — falta a peça deste"
            )
        if self.terceiros:
            return f"{self.terceiros} doc(s) de terceiros no processo, nenhum do escritório"
        return "nenhum documento do escritório no processo"


@dataclass
class AtoNaoCobrado:
    documento: DocumentoBenner
    ato_id: str
    ato_nome: str
    coluna: str
    coluna_presumida: bool
    valor: float | None


@dataclass
class ResultadoConferencia:
    competencia: str
    atos: list[ConferenciaAto]
    nao_cobrados: list[AtoNaoCobrado]
    docs_lidos: int
    docs_proprios: int

    def resumo(self) -> dict[str, Any]:
        por: dict[str, dict[str, float]] = {}
        for c in self.atos:
            s = por.setdefault(c.status, {"atos": 0, "valor": 0.0})
            s["atos"] += 1
            s["valor"] += c.auditada.valor_faturavel
        return {
            "competencia": self.competencia,
            "atos": len(self.atos),
            "por_status": {
                k: {"atos": int(v["atos"]), "valor": round(v["valor"], 2)}
                for k, v in sorted(por.items(), key=lambda kv: STATUS_ORDEM.index(kv[0]))
            },
            "nao_cobrados": len(self.nao_cobrados),
            "nao_cobrados_valor": round(sum(n.valor or 0 for n in self.nao_cobrados), 2),
            "docs_lidos": self.docs_lidos,
            "docs_proprios": self.docs_proprios,
        }


def _familia(a: str | None, b: str | None) -> bool:
    return bool(a and b and any(a in f and b in f for f in _FAMILIAS))


def _forca(linha_ato: str | None, d: DocumentoBenner) -> int:
    """0 = mesmo ato; 1 = genérico/protocolo/família; 2 = outro tipo."""
    if d.classe == "peca" and d.ato_id == linha_ato:
        return 0
    if d.classe in ("generico", "protocolo") or _familia(linha_ato, d.ato_id):
        return 1
    if linha_ato == "juntada_planilha_custas" and d.classe == "anexo":
        return 1
    return 2


def _competencia(c: str) -> tuple[date, date]:
    a, m = (int(x) for x in c.split("-"))
    ini = date(a, m, 1)
    fim = date(a + (m == 12), m % 12 + 1, 1) - timedelta(days=1)
    return ini, fim


def conferir(
    auditadas: list[LinhaAuditada],
    documentos: list[DocumentoBenner],
    competencia: str,
    proprios: Sequence[str] = PROPRIOS_PADRAO,
    cobrados_outras: Iterable[tuple[str, str | None, date | None]] = (),
    acao_por_cnj: dict[str, str] | None = None,
    marcar_achados: bool = True,
) -> ResultadoConferencia:
    """Confere a prestação contra os documentos do Benner.

    `cobrados_outras`: (cnj, ato_id, data) de outras competências, para não apontar como "não
    cobrado" um ato já faturado em outro mês. `acao_por_cnj`: tipo de ação conhecido por processo
    (prestações anteriores), para enquadrar a coluna dos não cobrados.
    """
    proprios = tuple(normalizar(p) for p in proprios)
    ini, fim = _competencia(competencia)
    meus: dict[str, list[DocumentoBenner]] = defaultdict(list)
    alheios: dict[str, list[DocumentoBenner]] = defaultdict(list)
    for d in documentos:
        alvo = meus if d.proprio(proprios) else alheios
        for proc in d.processos:
            alvo[proc].append(d)

    def janela(a: LinhaAuditada) -> tuple[date, date]:
        if a.linha.data:
            return a.linha.data - timedelta(days=ANTES_DIAS), a.linha.data + timedelta(
                days=DEPOIS_DIAS
            )
        return ini, fim + timedelta(days=DEPOIS_DIAS)

    def dist(a: LinhaAuditada, d: DocumentoBenner) -> int:
        if not d.data:
            return 999
        ref = a.linha.data or ini
        return abs((d.data - ref).days)

    usados: set[int] = set()
    resultado: dict[int, ConferenciaAto] = {}
    linhas = [a for a in auditadas if a.chave]
    for forca_max, status in ((0, "DOC_OK"), (1, "DOC_GENERICO"), (2, "DOC_DIVERGENTE")):
        for a in sorted(linhas, key=lambda x: (x.linha.data or ini, x.linha.origem)):
            if id(a) in resultado:
                continue
            j0, j1 = janela(a)
            k = _dig(a.linha.cnj)
            # do escritório em qualquer força; de terceiros só a peça do mesmo ato
            # (ex.: inicial que o preposto do Banco anexou ao abrir a ficha)
            cands = [(d, False) for d in meus.get(k, [])] + [
                (d, True) for d in alheios.get(k, []) if forca_max == 0
            ]
            cands = [
                (d, t)
                for d, t in cands
                if id(d) not in usados
                and d.data is not None
                and j0 <= d.data <= j1
                and _forca(a.ato_id, d) <= forca_max
            ]
            if not cands:
                continue
            d, terc = min(cands, key=lambda x: (_forca(a.ato_id, x[0]), x[1], dist(a, x[0])))
            usados.add(id(d))
            resultado[id(a)] = ConferenciaAto(a, status, d, terceiro=terc)

    atos: list[ConferenciaAto] = []
    for a in auditadas:
        c = resultado.get(id(a))
        if c is None:
            k = _dig(a.linha.cnj)
            livres = sorted(
                (d for d in meus.get(k, []) if id(d) not in usados), key=lambda x: dist(a, x)
            )
            st = "DOC_FORA_DA_JANELA" if livres else "SEM_DOCUMENTO"
            c = ConferenciaAto(a, st, None, livres[:3], len(alheios.get(k, [])))
            c.ja_usados = sum(1 for d in meus.get(k, []) if id(d) in usados)
        atos.append(c)
        if marcar_achados and c.status in ("SEM_DOCUMENTO", "DOC_DIVERGENTE", "DOC_FORA_DA_JANELA"):
            a.achados.append(Achado(c.status, "ALERTA", "Benner: " + c.detalhe))

    # direção inversa: peças do escritório sem ato cobrado
    tab = carregar_tabela()
    cobrados: dict[tuple[str, str], list[date]] = defaultdict(list)
    for a in auditadas:
        if a.ato_id:
            cobrados[(_dig(a.linha.cnj), a.ato_id)].append(a.linha.data or ini)
    for cnj, ato, dt in cobrados_outras:
        if ato:
            cobrados[(_dig(cnj), ato)].append(dt or ini)
    acoes = dict(acao_por_cnj or {})
    for a in auditadas:
        acoes.setdefault(_dig(a.linha.cnj), a.linha.acao)

    nao: list[AtoNaoCobrado] = []
    vistos: set[tuple[str, str, date]] = set()
    proprios_docs = {id(d): d for ds in meus.values() for d in ds}
    for d in sorted(proprios_docs.values(), key=lambda x: (x.data or ini, x.origem)):
        cnj = d.cnj_citado or d.cnj
        if (
            id(d) in usados
            or d.classe != "peca"
            or not d.ato_id
            or d.data is None
            or not (ini <= d.data <= fim)
            or nao_remunerado(d.nome)
        ):
            continue
        ato_tab = tab.atos.get(d.ato_id)
        if ato_tab is None:
            continue
        if ato_tab.resultado and not _SENTENCA.search(normalizar(d.nome + " " + d.tipo)):
            continue  # ato de resultado só com a decisão nos autos, não com o pedido
        if any(
            abs((d.data - x).days) <= JANELA_REPETICAO
            for proc in d.processos
            for x in cobrados.get((proc, d.ato_id), [])
        ):
            continue
        chave_nc = (cnj, d.ato_id, d.data)
        if chave_nc in vistos:
            continue
        vistos.add(chave_nc)
        acao = acoes.get(cnj) or acoes.get(d.cnj)
        col = enquadrar_coluna(acao, d.ficha, cnj)
        presumida = False
        if not acao and col == "demais":
            col, presumida = "rec_cred", True  # carteira BASA: recuperação de crédito
        valor = ato_tab.valor(col)
        if valor is None:
            continue  # N/A nessa coluna: não é cobrável
        nao.append(AtoNaoCobrado(d, d.ato_id, ato_tab.nome, col, presumida, valor))
        cobrados[(cnj, d.ato_id)].append(d.data)

    return ResultadoConferencia(
        competencia=competencia,
        atos=atos,
        nao_cobrados=nao,
        docs_lidos=len(documentos),
        docs_proprios=len(proprios_docs),
    )
