"""Serviço do M12 em produção: carteira BASA, prestações, DataJud e a rodada diária do motor.

Fluxo diário (worker): `sincronizar_datajud` (incremental) → `rodar_motor` (detecção +
conciliação com os atos cobrados) → oportunidades e relatório gravados no banco.

Todas as funções exigem o tenant no contexto (`current_tenant_id`), como o resto da plataforma.
"""

from __future__ import annotations

import io
import logging
import re
import uuid
from calendar import monthrange
from collections import defaultdict
from collections.abc import Iterable
from datetime import UTC, date, datetime, timedelta
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import load_workbook
from sqlalchemy import delete, func, select
from sqlalchemy.orm import Session

from app.core import cnj as cnjlib
from app.integrations.datajud.client import DataJudClient, alias_do_tribunal
from app.models.faturamento import (
    BasaAtoCobrado,
    BasaFicha,
    BasaRecebivel,
    DataJudProcesso,
    MotorExecucao,
    MotorOportunidade,
)
from app.modules.m12_faturamento import recebiveis as rcb
from app.modules.m12_faturamento.auditoria import FichaBenner, auditar, resumo
from app.modules.m12_faturamento.deteccao import Processo
from app.modules.m12_faturamento.planilhas import ler_atos
from app.modules.m12_faturamento.relatorio import gerar_relatorio, selecionar
from app.modules.m12_faturamento.tabela import normalizar, reconhecer_ato

log = logging.getLogger(__name__)

SITUACOES_ATIVAS = {"em andamento", "em encerramento", "em reativacao"}
JANELA_MOVIMENTOS_DIAS = 730  # o motor só olha os últimos 24 meses


def _dig(s: Any) -> str:
    return re.sub(r"\D", "", str(s or ""))


def _agora() -> datetime:
    return datetime.now(UTC)


_FUSO = ZoneInfo("America/Belem")


# --- carteira (Benner) -------------------------------------------------------------------


def importar_fichas(db: Session, tenant_id: uuid.UUID, arquivos: Iterable[bytes]) -> dict[str, int]:
    """Upsert das fichas a partir dos exports 'Fichas.xlsx' do Benner (cível e trabalhista)."""
    existentes = {f.ficha: f for f in db.scalars(select(BasaFicha))}
    novas = atualizadas = 0
    agora = _agora()
    for conteudo in arquivos:
        ws = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True).worksheets[0]
        rows = list(ws.iter_rows(values_only=True))
        cab = [normalizar(str(c or "")) for c in rows[0]]

        def col(nome: str, cab: list[str] = cab) -> int | None:
            return next((i for i, c in enumerate(cab) if c.startswith(nome)), None)

        i_f, i_s = cab.index("ficha"), cab.index("situacao")
        i_c = next(i for i, c in enumerate(cab) if c.startswith("n") and "processo" in c)
        i_a, i_v, i_d = col("adverso"), col("valor da causa"), col("cadastrado")
        for r in rows[1:]:
            if not r[i_f]:
                continue
            ficha = str(r[i_f]).strip()[:30]
            f = existentes.get(ficha)
            if f is None:
                f = BasaFicha(tenant_id=tenant_id, ficha=ficha, cnj="", situacao="")
                db.add(f)
                existentes[ficha] = f
                novas += 1
            else:
                atualizadas += 1
            f.cnj = (_dig(r[i_c]) or str(r[i_c] or ""))[:25]
            f.situacao = str(r[i_s] or "")[:40]
            f.categoria = "trabalhista" if ficha.upper().startswith("TRAB") else "civel"
            f.adverso = str(r[i_a])[:200] if i_a is not None and r[i_a] else None
            v = r[i_v] if i_v is not None else None
            f.valor_causa = float(v) if isinstance(v, int | float) else None
            d = r[i_d] if i_d is not None else None
            f.cadastrado_em = d if isinstance(d, datetime) else None
            f.importado_em = agora
    db.commit()
    total = db.scalar(select(func.count()).select_from(BasaFicha)) or 0
    return {"novas": novas, "atualizadas": atualizadas, "total": int(total)}


def fichas_ativas(db: Session) -> list[BasaFicha]:
    return [
        f
        for f in db.scalars(select(BasaFicha))
        if normalizar(f.situacao) in SITUACOES_ATIVAS and len(f.cnj) == 20
    ]


# --- prestações de contas ----------------------------------------------------------------


def importar_prestacao(
    db: Session, tenant_id: uuid.UUID, competencia: str, conteudo: bytes
) -> dict[str, Any]:
    """Substitui a competência inteira, atualiza os recebíveis por UF e devolve a pré-auditoria."""
    linhas = ler_atos(io.BytesIO(conteudo))  # type: ignore[arg-type]
    benner = {
        f.ficha: FichaBenner(f.ficha, f.cnj, f.situacao) for f in db.scalars(select(BasaFicha))
    }
    auditadas = auditar(linhas, competencia, benner or None)  # normaliza a UF
    db.execute(delete(BasaAtoCobrado).where(BasaAtoCobrado.competencia == competencia))
    for a in auditadas:
        ln = a.linha
        db.add(
            BasaAtoCobrado(
                tenant_id=tenant_id,
                competencia=competencia,
                cnj=(_dig(ln.cnj) or ln.cnj)[:25],
                ficha=(ln.ficha or None) and ln.ficha[:30],
                uf=(ln.uf or None) and ln.uf[:2],
                ato_id=a.ato_id or reconhecer_ato(ln.ato_texto),
                ato_texto=ln.ato_texto,
                data=ln.data,
                valor=ln.valor,
            )
        )
    _recebiveis_da_prestacao(db, tenant_id, competencia, auditadas)
    db.commit()
    return {
        "competencia": competencia,
        "linhas": len(linhas),
        "pre_auditoria": resumo(auditadas),
    }


def _recebiveis_da_prestacao(
    db: Session, tenant_id: uuid.UUID, competencia: str, auditadas: list[Any]
) -> None:
    """Cria/atualiza uma linha de recebível por UF; preserva o que a equipe lançou à mão."""
    tot: dict[str, dict[str, float]] = defaultdict(lambda: {"dec": 0.0, "fat": 0.0, "n": 0})
    for a in auditadas:
        t = tot[rcb.grupo_uf(a.linha.uf)]
        t["dec"] += float(a.linha.valor or 0)
        t["fat"] += a.valor_faturavel
        t["n"] += 1
    existentes = {
        r.uf: r
        for r in db.scalars(select(BasaRecebivel).where(BasaRecebivel.competencia == competencia))
    }
    for uf, t in tot.items():
        r = existentes.get(uf)
        if r is None:
            r = BasaRecebivel(tenant_id=tenant_id, competencia=competencia, uf=uf, historico=[])
            db.add(r)
        r.valor_declarado = round(t["dec"], 2)
        r.valor_faturavel = round(t["fat"], 2)
        r.atos = int(t["n"])


def cobertura_das_prestacoes(db: Session) -> tuple[date, date] | None:
    comps = sorted(set(db.scalars(select(BasaAtoCobrado.competencia))))
    if not comps:
        return None
    a0, m0 = (int(x) for x in comps[0].split("-"))
    a1, m1 = (int(x) for x in comps[-1].split("-"))
    return date(a0, m0, 1), date(a1, m1, monthrange(a1, m1)[1])


def conferir_prestacao(
    db: Session, competencia: str, prestacao: bytes, documentos: Iterable[bytes]
) -> tuple[bytes, dict[str, Any]]:
    """Pré-auditoria + conferente do Benner, sem gravar nada: o escritório roda antes de enviar.

    Usa a carteira já importada (fichas) e as outras competências gravadas, para não apontar
    como "não cobrado" um ato faturado em outro mês.
    """
    import tempfile
    from pathlib import Path

    from app.modules.m12_faturamento.conferencia import conferir, ler_documentos
    from app.modules.m12_faturamento.planilhas import gerar_excel

    linhas = ler_atos(io.BytesIO(prestacao))  # type: ignore[arg-type]
    benner = {
        f.ficha: FichaBenner(f.ficha, f.cnj, f.situacao) for f in db.scalars(select(BasaFicha))
    }
    auditadas = auditar(linhas, competencia, benner or None)
    outras = [
        (a.cnj, a.ato_id, a.data)
        for a in db.scalars(select(BasaAtoCobrado).where(BasaAtoCobrado.competencia != competencia))
    ]
    res = conferir(auditadas, ler_documentos(list(documentos)), competencia, cobrados_outras=outras)
    with tempfile.TemporaryDirectory() as d:
        out = Path(d) / "conf.xlsx"
        rs = gerar_excel(auditadas, competencia, out, conferencia=res)
        return out.read_bytes(), rs


# --- recebíveis --------------------------------------------------------------------------

CAMPOS_RECEBIVEL = {
    "enviado_em": "data",
    "conferido_em": "data",
    "valor_aprovado": "valor",
    "nf_numero": "texto",
    "nf_emitida_em": "data",
    "nf_valor": "valor",
    "pago_em": "data",
    "valor_pago": "valor",
    "observacao": "texto",
}


def _rec(r: BasaRecebivel) -> rcb.Recebivel:
    return rcb.Recebivel(
        competencia=r.competencia,
        uf=r.uf,
        valor_faturavel=r.valor_faturavel,
        valor_declarado=r.valor_declarado,
        enviado_em=r.enviado_em,
        conferido_em=r.conferido_em,
        valor_aprovado=r.valor_aprovado,
        nf_numero=r.nf_numero,
        nf_emitida_em=r.nf_emitida_em,
        nf_valor=r.nf_valor,
        pago_em=r.pago_em,
        valor_pago=r.valor_pago,
    )


def recebivel_dict(r: BasaRecebivel, hoje: date) -> dict[str, Any]:
    out: dict[str, Any] = {
        "id": str(r.id),
        "competencia": r.competencia,
        "uf": r.uf,
        "atos": r.atos,
        "valor_declarado": r.valor_declarado,
        "valor_faturavel": r.valor_faturavel,
        "historico": r.historico or [],
    }
    for c, tipo in CAMPOS_RECEBIVEL.items():
        v = getattr(r, c)
        out[c] = v.isoformat() if tipo == "data" and v else v
    out.update(rcb.situacao(_rec(r), hoje).as_dict())
    return out


def listar_recebiveis(db: Session, hoje: date | None = None) -> dict[str, Any]:
    hoje = hoje or _agora().astimezone(_FUSO).date()
    itens = [
        recebivel_dict(r, hoje)
        for r in db.scalars(
            select(BasaRecebivel).order_by(BasaRecebivel.competencia.desc(), BasaRecebivel.uf)
        )
    ]
    em_aberto = sum(i["em_aberto"] for i in itens)
    vencido = sum(
        i["em_aberto"] for i in itens if any(a["codigo"] == "VENCIDO" for a in i["alertas"])
    )
    return {
        "hoje": hoje.isoformat(),
        "em_aberto": round(em_aberto, 2),
        "vencido": round(vencido, 2),
        "alertas": sum(len(i["alertas"]) for i in itens),
        "itens": itens,
    }


def atualizar_recebivel(
    db: Session, rid: uuid.UUID, campos: dict[str, Any], por: str, hoje: date | None = None
) -> dict[str, Any]:
    r = db.get(BasaRecebivel, rid)
    if r is None:
        raise KeyError(rid)
    hist = list(r.historico or [])
    for c, v in campos.items():
        tipo = CAMPOS_RECEBIVEL.get(c)
        if tipo is None:
            raise ValueError(f"campo não editável: {c}")
        if v in ("", None):
            novo: Any = None
        elif tipo == "data":
            novo = date.fromisoformat(str(v)[:10])
        elif tipo == "valor":
            novo = round(float(v), 2)
        else:
            novo = str(v)[:2000] if c == "observacao" else str(v)[:40]
        antigo = getattr(r, c)
        if antigo == novo:
            continue
        setattr(r, c, novo)
        hist.append(
            {
                "em": _agora().isoformat(timespec="seconds"),
                "por": por[:80],
                "campo": c,
                "de": antigo.isoformat() if isinstance(antigo, date) else antigo,
                "para": novo.isoformat() if isinstance(novo, date) else novo,
            }
        )
    r.historico = hist
    db.commit()
    return recebivel_dict(r, hoje or _agora().astimezone(_FUSO).date())


# --- DataJud -----------------------------------------------------------------------------


def _enxugar(source: dict[str, Any], limite: date) -> dict[str, Any]:
    movs = []
    for m in source.get("movimentos") or []:
        d = (m.get("dataHora") or "")[:10]
        if d and d >= limite.isoformat():
            movs.append(
                [
                    m.get("codigo"),
                    m.get("nome"),
                    d,
                    [
                        c.get("nome")
                        for c in (m.get("complementosTabelados") or [])
                        if c.get("nome")
                    ],
                ]
            )
    return {
        "grau": source.get("grau"),
        "classe": (source.get("classe") or {}).get("nome"),
        "atualizado": source.get("dataHoraUltimaAtualizacao"),
        "movimentos": movs,
        "ultima": max(
            ((m.get("dataHora") or "")[:10] for m in source.get("movimentos") or []), default=None
        ),
    }


def sincronizar_datajud(
    db: Session,
    tenant_id: uuid.UUID,
    client: DataJudClient,
    completo: bool = False,
    hoje: date | None = None,
) -> dict[str, Any]:
    """Busca no DataJud os processos ativos da carteira. Incremental por padrão: só o que o
    CNJ atualizou desde a busca anterior (menos 2 dias de folga); processo nunca buscado vai
    sempre."""
    hoje = hoje or date.today()
    limite = hoje - timedelta(days=JANELA_MOVIMENTOS_DIAS)
    salvos = {p.cnj: p for p in db.scalars(select(DataJudProcesso))}
    ultima_busca = None if completo else max((p.buscado_em for p in salvos.values()), default=None)
    desde = ultima_busca - timedelta(days=2) if ultima_busca else None
    por_tribunal: dict[str, list[str]] = defaultdict(list)
    novos_por_tribunal: dict[str, list[str]] = defaultdict(list)
    for f in fichas_ativas(db):
        t = alias_do_tribunal(f.cnj)
        if not t:
            continue
        (por_tribunal if f.cnj in salvos else novos_por_tribunal)[t].append(f.cnj)
    agora = _agora()
    stats: dict[str, Any] = {
        "desde": desde.isoformat() if desde else None,
        "hits": 0,
        "processos_atualizados": 0,
        "por_tribunal": {},
    }
    tocados: set[str] = set()

    def gravar(hits: Iterable[dict[str, Any]]) -> int:
        n = 0
        for h in hits:
            s = h.get("_source") or {}
            num = s.get("numeroProcesso")
            if not num:
                continue
            p = salvos.get(num)
            if p is None:
                p = DataJudProcesso(
                    tenant_id=tenant_id,
                    cnj=num,
                    tribunal=s.get("tribunal") or "",
                    graus={},
                    buscado_em=agora,
                )
                db.add(p)
                salvos[num] = p
            graus = dict(p.graus or {})
            graus[str(h.get("_id"))] = _enxugar(s, limite)
            p.graus = graus
            classes = [g["classe"] for g in graus.values() if g.get("grau") in ("G1", "JE")]
            p.classe = (classes or [g.get("classe") for g in graus.values()] or [None])[0]
            ult = [g["ultima"] for g in graus.values() if g.get("ultima")]
            p.ultima_movimentacao = date.fromisoformat(max(ult)) if ult else None
            atu = [g["atualizado"] for g in graus.values() if g.get("atualizado")]
            if atu:
                p.datajud_atualizado = datetime.fromisoformat(max(atu).replace("Z", "+00:00"))
            tocados.add(num)
            n += 1
        return n

    for t in sorted(set(por_tribunal) | set(novos_por_tribunal)):
        n = gravar(client.buscar_por_numeros(t, novos_por_tribunal.get(t, [])))
        n += gravar(client.buscar_por_numeros(t, por_tribunal.get(t, []), desde=desde))
        stats["por_tribunal"][t] = n
        stats["hits"] += n
        db.commit()
    for p in salvos.values():
        p.buscado_em = agora
    db.commit()
    stats["processos_atualizados"] = len(tocados)
    return stats


def _processo(p: DataJudProcesso) -> Processo:
    movs: list[dict[str, Any]] = []
    for g in (p.graus or {}).values():
        for cod, nome, d, comps in g.get("movimentos") or []:
            movs.append(
                {
                    "codigo": cod,
                    "nome": nome,
                    "dataHora": f"{d}T00:00:00",
                    "complementosTabelados": [{"nome": c} for c in comps],
                }
            )
    proc = Processo(p.cnj, p.classe or "", p.tribunal, movs)
    proc.ultima_movimentacao = p.ultima_movimentacao
    return proc


# --- rodada do motor ---------------------------------------------------------------------


def rodar_motor(db: Session, tenant_id: uuid.UUID, hoje: date | None = None) -> MotorExecucao:
    hoje = hoje or date.today()
    ex = MotorExecucao(tenant_id=tenant_id, iniciado_em=_agora(), referencia=hoje)
    db.add(ex)
    db.commit()
    try:
        ativas = {f.cnj: f.ficha for f in fichas_ativas(db)}
        processos = [_processo(p) for p in db.scalars(select(DataJudProcesso)) if p.cnj in ativas]
        cobrados: dict[str, list[tuple[date, str]]] = defaultdict(list)
        for a in db.scalars(select(BasaAtoCobrado)):
            if a.ato_id and a.data:
                cobrados[a.cnj].append((a.data, a.ato_id))
        cob = cobertura_das_prestacoes(db) or (hoje - timedelta(days=90), hoje - timedelta(days=90))
        ops = selecionar(processos, cobrados, hoje, cob)
        xlsx, res = gerar_relatorio(ops, len(processos), ativas, hoje, cob)
        db.execute(delete(MotorOportunidade))  # só a última rodada fica em tabela
        for o in ops:
            db.add(
                MotorOportunidade(
                    tenant_id=tenant_id,
                    execucao_id=ex.id,
                    cnj=o.cnj,
                    ficha=ativas.get(o.cnj),
                    evento=o.evento,
                    data_evento=o.data_evento,
                    ato_id=o.ato_id,
                    valor=o.valor,
                    papel=o.papel,
                    prazo=o.prazo,
                    confianca=o.confianca,
                    movimento=o.movimento[:2000] if o.movimento else None,
                    observacao=o.observacao,
                    status=o.status,
                    atendida_por=o.atendida_por,
                )
            )
        ex.estatisticas = {
            **{k: v for k, v in res.items() if k != "por_evento"},
            "carteira_ativa": len(ativas),
            "cobertura": [cob[0].isoformat(), cob[1].isoformat()],
            "por_evento": res["por_evento"],
        }
        ex.relatorio_xlsx = xlsx
        ex.status = "ok"
    except Exception as e:  # registra e propaga para o log do worker
        db.rollback()
        ex.status, ex.erro = "erro", f"{type(e).__name__}: {e}"[:4000]
        raise
    finally:
        ex.terminado_em = _agora()
        db.add(ex)
        db.commit()
    return ex


def cnj_formatado(d: str) -> str:
    return cnjlib.normalize(d) or d
