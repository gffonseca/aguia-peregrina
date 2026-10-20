"""CLI do motor de oportunidades.

python -m app.modules.m12_faturamento.oportunidades_cli \\
    --movimentos raw.jsonl --prestacoes jul.xlsx ago.xlsx set.xlsx \\
    --benner Fichas.xlsx ... --hoje 2026-10-20 --cobertura 2026-07-01:2026-09-30 --saida out.xlsx

`--cobertura` = período coberto pelas prestações informadas. Eventos cujo prazo cai dentro
dele são conciliados (ATENDIDA/PERDIDA); eventos posteriores entram na fila (ABERTA).
"""

from __future__ import annotations

import argparse
import json
import re
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from app.modules.m12_faturamento.deteccao import (
    DIAS_PARA_ATENDER,
    Oportunidade,
    conciliar,
    detectar,
    unir_graus,
)
from app.modules.m12_faturamento.planilhas import ler_atos, ler_benner
from app.modules.m12_faturamento.tabela import carregar_tabela, reconhecer_ato


def _d(s: str) -> date:
    return datetime.strptime(s, "%Y-%m-%d").date()


def _dig(s: str | None) -> str:
    return re.sub(r"\D", "", s or "")


def executar(
    movimentos: Path,
    prestacoes: list[Path],
    benner: list[Path],
    hoje: date,
    cobertura: tuple[date, date],
    saida: Path,
) -> dict[str, Any]:
    docs = [json.loads(linha) for linha in movimentos.read_text(encoding="utf-8").splitlines()]
    processos = unir_graus(docs)
    fichas = {_dig(f.cnj): f.ficha for f in ler_benner(benner).values()} if benner else {}

    cobrados: dict[str, list[tuple[date, str]]] = defaultdict(list)
    for p in prestacoes:
        for ln in ler_atos(p):
            a = reconhecer_ato(ln.ato_texto)
            if a and ln.data:
                cobrados[_dig(ln.cnj)].append((ln.data, a))

    ini, fim = cobertura
    todas: list[Oportunidade] = []
    for proc in processos.values():
        ops = detectar(proc, hoje, desde=ini - timedelta(days=DIAS_PARA_ATENDER))
        conciliar(ops, cobrados, hoje)
        for o in ops:
            if o.status == "ALERTA" or o.evento == "processo_sem_evento_90_dias":
                pass
            elif o.data_evento < ini - timedelta(days=3) and o.status != "ATENDIDA":
                continue  # evento anterior à cobertura e não atendido: sem como julgar
            elif (
                o.status == "PERDIDA"
                and (o.prazo or o.data_evento + timedelta(days=DIAS_PARA_ATENDER)) > fim
            ):
                o.status = "ABERTA"  # janela além das prestações informadas: ainda na fila
            todas.append(o)

    resumo = _escrever(todas, processos, fichas, hoje, cobertura, saida)
    return resumo


def _escrever(
    ops: list[Oportunidade],
    processos: dict[str, Any],
    fichas: dict[str, str],
    hoje: date,
    cobertura: tuple[date, date],
    saida: Path,
) -> dict[str, Any]:
    tab = carregar_tabela()
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Resumo"
    bold = Font(bold=True)
    ws.append(
        [
            f"Motor de oportunidades BASA — {hoje:%d/%m/%Y} — cobertura das prestações "
            f"{cobertura[0]:%d/%m/%Y} a {cobertura[1]:%d/%m/%Y}"
        ]
    )
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([f"Processos lidos no DataJud: {len(processos)}"])
    ws.append([])
    ws.append(
        [
            "Evento",
            "Ato",
            "Papel da IA",
            "Confiança",
            "Atendidas",
            "Perdidas",
            "Abertas (fila)",
            "R$ perdido",
            "R$ em aberto",
        ]
    )
    for c in ws[4]:
        c.font = bold
    agg: dict[tuple[str, str, str, str], Counter[str]] = defaultdict(Counter)
    val: dict[tuple[str, str, str, str], defaultdict[str, float]] = defaultdict(
        lambda: defaultdict(float)
    )
    for o in ops:
        if o.status == "ALERTA":
            continue
        k = (o.evento, o.ato_id or "", o.papel or "", o.confianca)
        agg[k][o.status] += 1
        val[k][o.status] += o.valor or 0
    linhas = sorted(agg, key=lambda k: -(val[k]["PERDIDA"] + val[k]["ABERTA"]))
    tot: defaultdict[str, float] = defaultdict(float)
    for k in linhas:
        nome_ato = tab.atos[k[1]].nome if k[1] in tab.atos else k[1]
        ws.append(
            [
                k[0],
                nome_ato,
                k[2],
                k[3],
                agg[k]["ATENDIDA"],
                agg[k]["PERDIDA"],
                agg[k]["ABERTA"],
                round(val[k]["PERDIDA"], 2),
                round(val[k]["ABERTA"], 2),
            ]
        )
        for s in ("ATENDIDA", "PERDIDA", "ABERTA"):
            tot[s] += agg[k][s]
            tot["v" + s] += val[k][s]
    ws.append(
        [
            "TOTAL",
            "",
            "",
            "",
            tot["ATENDIDA"],
            tot["PERDIDA"],
            tot["ABERTA"],
            round(tot["vPERDIDA"], 2),
            round(tot["vABERTA"], 2),
        ]
    )
    for c in ws[ws.max_row]:
        c.font = bold
    for i, w in enumerate([40, 46, 10, 10, 10, 10, 13, 13, 13], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    def aba(nome: str, filtro: Any) -> None:
        sh = wb.create_sheet(nome)
        cab = [
            "Processo",
            "Ficha",
            "Evento",
            "Data do evento",
            "Ato cabível",
            "Valor",
            "Papel da IA",
            "Prazo de referência",
            "Confiança",
            "Movimento",
            "Observação",
            "Atendida por",
        ]
        sh.append(cab)
        for c in sh[1]:
            c.font = bold
        for o in sorted(filter(filtro, ops), key=lambda x: (-(x.valor or 0), x.data_evento)):
            cnj = o.cnj
            fmt = f"{cnj[:7]}-{cnj[7:9]}.{cnj[9:13]}.{cnj[13]}.{cnj[14:16]}.{cnj[16:]}"
            sh.append(
                [
                    fmt,
                    fichas.get(cnj, ""),
                    o.evento,
                    o.data_evento,
                    tab.atos[o.ato_id].nome if o.ato_id in tab.atos else "",
                    o.valor,
                    o.papel,
                    o.prazo,
                    o.confianca,
                    o.movimento,
                    o.observacao,
                    o.atendida_por or "",
                ]
            )
        for i, w in enumerate([26, 18, 34, 12, 40, 9, 10, 12, 10, 40, 46, 26], start=1):
            sh.column_dimensions[get_column_letter(i)].width = w
        for row in sh.iter_rows(min_row=2):
            row[3].number_format = "DD/MM/YYYY"
            row[7].number_format = "DD/MM/YYYY"
        sh.freeze_panes = "A2"

    aba("Fila aberta", lambda o: o.status == "ABERTA" and o.evento != "processo_sem_evento_90_dias")
    aba("Perdidas", lambda o: o.status == "PERDIDA")
    aba("Sem movimento 90 dias", lambda o: o.evento == "processo_sem_evento_90_dias")
    aba("Alertas", lambda o: o.status == "ALERTA" and o.data_evento >= hoje - timedelta(days=120))
    aba("Atendidas", lambda o: o.status == "ATENDIDA")
    saida.parent.mkdir(parents=True, exist_ok=True)
    wb.save(saida)
    alertas = Counter(o.evento for o in ops if o.status == "ALERTA")
    return {
        "processos": len(processos),
        "atendidas": tot["ATENDIDA"],
        "perdidas": tot["PERDIDA"],
        "abertas": tot["ABERTA"],
        "valor_perdido": round(tot["vPERDIDA"], 2),
        "valor_aberto": round(tot["vABERTA"], 2),
        "alertas": dict(alertas),
        "por_evento": {
            k[0] + "|" + k[3]: {
                **agg[k],
                "R$perdido": round(val[k]["PERDIDA"]),
                "R$aberto": round(val[k]["ABERTA"]),
            }
            for k in linhas
        },
    }


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--movimentos", required=True, type=Path)
    ap.add_argument("--prestacoes", nargs="*", default=[], type=Path)
    ap.add_argument("--benner", nargs="*", default=[], type=Path)
    ap.add_argument("--hoje", default=date.today().isoformat())
    ap.add_argument("--cobertura", required=True, help="AAAA-MM-DD:AAAA-MM-DD")
    ap.add_argument("--saida", required=True, type=Path)
    a = ap.parse_args()
    ini, fim = (_d(x) for x in a.cobertura.split(":"))
    r = executar(a.movimentos, a.prestacoes, a.benner, _d(a.hoje), (ini, fim), a.saida)
    print(json.dumps(r, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
