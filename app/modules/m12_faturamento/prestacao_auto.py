"""Prestação de contas montada a partir das evidências (Fase 2).

Em vez de a equipe listar à mão o que fez no mês, o sistema propõe as linhas:
1. Peças do escritório no Benner com data no mês, reconhecidas como ato da tabela (mesma
   regra do conferente), sem ato igual já cobrado no mesmo processo em ±45 dias;
2. Atos de resultado detectados pelo motor no tribunal no mês (citação efetivada, penhora,
   extinção...) ainda não cobrados — marcados "conferir", porque dependem de confirmar o
   resultado nos autos.

A saída é a planilha no layout que o escritório já usa (NPC, Nº Processo, Comarca, UF, Ação,
Parte Adversa, Ato processual realizado, Data do protocolo, Valor) + colunas de evidência. A
equipe revisa, apaga o que não procede e sobe em POST /faturamento/basa/prestacoes, que roda a
pré-auditoria de sempre. Nada vai ao Banco sem esse passo humano.
"""

from __future__ import annotations

import hashlib
import io
import uuid
from datetime import UTC, date, datetime, timedelta
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models.benner import BennerDocumento
from app.models.faturamento import (
    BasaAtoCobrado,
    BasaFicha,
    DataJudProcesso,
    MotorOportunidade,
)
from app.modules.m12_faturamento.conferencia import (
    DocumentoBenner,
    _competencia,
    cnj_fmt,
    conferir,
)
from app.modules.m12_faturamento.tabela import carregar_tabela, enquadrar_coluna
from app.modules.m13_triagem.classificar import uf_do_cnj

UF_NOME = {"PA": "Pará", "AM": "Amazonas"}
CABECALHO = [
    "NPC",
    "Nº Processo",
    "Comarca",
    "UF",
    "Ação",
    "Parte Adversa",
    "Ato processual realizado",
    "Data do protocolo",
    "Valor",
    "Origem",
    "Evidência",
    "Conferir",
]


def _chave_doc(d: DocumentoBenner) -> str:
    raw = "|".join(
        [d.cnj, d.ficha, d.data.isoformat() if d.data else "", d.nome, d.arquivo, d.incluido_por]
    )
    return hashlib.sha1(raw.encode()).hexdigest()


def guardar_documentos(
    db: Session, tenant_id: uuid.UUID, docs: list[DocumentoBenner]
) -> dict[str, int]:
    ja = set(db.scalars(select(BennerDocumento.chave)))
    agora = datetime.now(UTC)
    novos = 0
    for d in docs:
        k = _chave_doc(d)
        if k in ja:
            continue
        ja.add(k)
        db.add(
            BennerDocumento(
                tenant_id=tenant_id,
                chave=k,
                data=d.data,
                nome=d.nome[:500],
                tipo=(d.tipo or None) and d.tipo[:200],
                ficha=(d.ficha or None) and d.ficha[:30],
                cnj=(d.cnj or None) and d.cnj[:25],
                incluido_por=(d.incluido_por or None) and d.incluido_por[:200],
                arquivo=(d.arquivo or None) and d.arquivo[:500],
                importado_em=agora,
            )
        )
        novos += 1
    db.commit()
    total = len(ja)
    return {"novos": novos, "total": total}


def documentos_guardados(
    db: Session, de: date | None = None, ate: date | None = None
) -> list[DocumentoBenner]:
    q = select(BennerDocumento)
    if de:
        q = q.where(BennerDocumento.data >= de)
    if ate:
        q = q.where(BennerDocumento.data <= ate)
    return [
        DocumentoBenner(
            data=r.data,
            nome=r.nome,
            tipo=r.tipo or "",
            ficha=r.ficha or "",
            cnj=r.cnj or "",
            incluido_por=r.incluido_por or "",
            arquivo=r.arquivo or "",
        )
        for r in db.scalars(q)
    ]


def _contexto(db: Session) -> tuple[dict[str, dict[str, str]], dict[str, BasaFicha]]:
    """Comarca/ação/parte por processo, das prestações anteriores (a mais recente vence)."""
    ctx: dict[str, dict[str, str]] = {}
    for a in db.scalars(select(BasaAtoCobrado).order_by(BasaAtoCobrado.competencia)):
        c = ctx.setdefault(a.cnj, {})
        for k in ("comarca", "acao", "parte_adversa"):
            v = getattr(a, k, None)
            if v:
                c[k] = v
    fichas = {f.cnj: f for f in db.scalars(select(BasaFicha))}
    return ctx, fichas


def montar(db: Session, competencia: str) -> tuple[bytes, dict[str, Any]]:
    ini, fim = _competencia(competencia)
    docs = documentos_guardados(db, ini - timedelta(days=60), fim)
    cobrados = [(a.cnj, a.ato_id, a.data) for a in db.scalars(select(BasaAtoCobrado))]
    ctx, fichas = _contexto(db)
    classes = {p.cnj: p.classe for p in db.scalars(select(DataJudProcesso))}
    acao_por_cnj = {k: v["acao"] for k, v in ctx.items() if v.get("acao")}
    for k, v in classes.items():
        if v:
            acao_por_cnj.setdefault(k, v)

    res = conferir([], docs, competencia, cobrados_outras=cobrados, acao_por_cnj=acao_por_cnj)
    linhas: list[dict[str, Any]] = []
    vistos: set[tuple[str, str, date]] = set()
    for n in res.nao_cobrados:
        d = n.documento
        cnj = d.cnj_citado or d.cnj
        assert d.data is not None
        vistos.add((cnj, n.ato_id, d.data))
        linhas.append(
            {
                "cnj": cnj,
                "ficha": d.ficha,
                "ato": n.ato_nome,
                "data": d.data,
                "valor": n.valor,
                "origem": "peça no Benner"
                + (
                    " (tipo de ação presumido: recuperação de crédito)"
                    if n.coluna_presumida
                    else ""
                ),
                "evidencia": f"{d.nome or d.tipo} [{d.arquivo}]",
                "conferir": "peça cita outro processo" if d.cnj_citado else "",
            }
        )

    # atos de resultado detectados no tribunal no mês e ainda não cobrados
    tab = carregar_tabela()
    ja = {(c, a) for c, a, _ in cobrados}
    for o in db.scalars(select(MotorOportunidade).where(MotorOportunidade.status == "ABERTA")):
        ato = tab.atos.get(o.ato_id or "")
        if ato is None or not ato.resultado or not (ini <= o.data_evento <= fim):
            continue
        if (o.cnj, o.ato_id) in ja or (o.cnj, o.ato_id, o.data_evento) in vistos:
            continue
        f = fichas.get(o.cnj)
        col = enquadrar_coluna(acao_por_cnj.get(o.cnj), f.ficha if f else None, o.cnj)
        valor = ato.valor(col) or ato.valor("rec_cred")
        if valor is None:
            continue
        linhas.append(
            {
                "cnj": o.cnj,
                "ficha": f.ficha if f else (o.ficha or ""),
                "ato": ato.nome,
                "data": o.data_evento,
                "valor": valor,
                "origem": "resultado no tribunal (DataJud)",
                "evidencia": f"{o.movimento or o.evento} em {o.data_evento:%d/%m/%Y}",
                "conferir": "confirmar o resultado nos autos",
            }
        )

    linhas.sort(key=lambda x: (uf_do_cnj(x["cnj"]) or "ZZ", x["data"], x["cnj"]))
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = f"Prestação {competencia}"
    ws.append(CABECALHO)
    for c in ws[1]:
        c.font = Font(bold=True)
    amarelo = PatternFill("solid", fgColor="FFF2CC")
    for x in linhas:
        c = ctx.get(x["cnj"], {})
        f = fichas.get(x["cnj"])
        uf = uf_do_cnj(x["cnj"])
        ws.append(
            [
                x["ficha"],
                cnj_fmt(x["cnj"]),
                c.get("comarca", ""),
                UF_NOME.get(uf or "", uf or ""),
                c.get("acao") or classes.get(x["cnj"]) or "",
                c.get("parte_adversa") or (f.adverso if f else "") or "",
                x["ato"],
                x["data"],
                x["valor"],
                x["origem"],
                x["evidencia"],
                x["conferir"],
            ]
        )
        ws.cell(row=ws.max_row, column=8).number_format = "DD/MM/YYYY"
        ws.cell(row=ws.max_row, column=9).number_format = "#,##0.00"
        if x["conferir"]:
            ws.cell(row=ws.max_row, column=12).fill = amarelo
    for col, w in zip(
        "ABCDEFGHIJKL", (18, 26, 18, 10, 30, 34, 50, 12, 10, 26, 60, 26), strict=True
    ):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    por_uf: dict[str, dict[str, float]] = {}
    for x in linhas:
        u = por_uf.setdefault(uf_do_cnj(x["cnj"]) or "outras", {"atos": 0, "valor": 0.0})
        u["atos"] += 1
        u["valor"] += x["valor"] or 0
    return buf.getvalue(), {
        "competencia": competencia,
        "linhas": len(linhas),
        "valor": round(sum(x["valor"] or 0 for x in linhas), 2),
        "a_conferir": sum(1 for x in linhas if x["conferir"]),
        "por_uf": {
            k: {"atos": int(v["atos"]), "valor": round(v["valor"], 2)} for k, v in por_uf.items()
        },
        "documentos_considerados": len(docs),
    }
