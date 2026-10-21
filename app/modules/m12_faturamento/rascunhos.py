"""IA redige a peça a partir de um item da fila (Fase 2).

Entrada da IA, só o que o sistema sabe de verdade sobre o processo:
- a oportunidade (evento, data, movimento do tribunal, ato da tabela);
- a ficha do Benner (número, adverso, valor da causa) e a classe/últimos movimentos (DataJud);
- comarca/vara e tipo de ação das prestações anteriores;
- o TEXTO INTEGRAL da intimação mais recente do processo no DJEN, quando houver — é o que
  permite à IA ler a sentença/decisão de verdade e redigir a apelação ou a manifestação.

Regras: a peça é RASCUNHO para revisão de advogado; nenhum fato que não esteja nas fontes;
nenhuma jurisprudência citada por número (só dispositivos de lei, e só os da lista abaixo);
lacunas como «[REVISOR: …]». A saída é Markdown + .docx.
"""

from __future__ import annotations

import io
import re
import uuid
from datetime import UTC, date, datetime
from typing import Any

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models.faturamento import (
    BasaAtoCobrado,
    BasaFicha,
    DataJudProcesso,
    MotorOportunidade,
    RascunhoPeca,
)
from app.models.triagem import TriagemEmail
from app.modules.m12_faturamento.fila import ATO, ROTULO
from app.modules.m12_faturamento.tabela import carregar_tabela

PROMPT_VERSAO = "rascunho-basa-v1"

INSTRUCOES_POR_ATO: dict[str, str] = {
    "apelacao": (
        "Apelação (CPC, arts. 1.009 a 1.014) contra a sentença descrita nas fontes. Use os "
        "fundamentos da própria sentença (copiados do texto da intimação) para atacá-los "
        "ponto a ponto. Se as fontes não trouxerem o teor da sentença, redija só a estrutura "
        "(endereçamento, tempestividade, preparo, razões com «[REVISOR: inserir fundamentos "
        "da sentença]», pedidos)."
    ),
    "contrarrazoes": "Contrarrazões ao recurso da parte adversa (CPC, art. 1.010, § 1º).",
    "embargos_declaracao": "Embargos de declaração (CPC, arts. 1.022 a 1.026), só se houver "
    "omissão/contradição/obscuridade identificável nas fontes; senão, estrutura com lacunas.",
    "peticao_interlocutoria": (
        "Petição interlocutória que peça providência NOVA e útil ao andamento (ex.: pesquisa de "
        "bens via SISBAJUD/RENAJUD/INFOJUD, citação por outro meio, penhora, expedição de "
        "mandado). Proibido apenas reiterar pedido anterior: o contrato não remunera reiteração. "
        "Se o evento for prescrição intercorrente, manifeste-se contra ela (CPC, art. 921, §§ 4º "
        "e 5º) indicando diligências realizadas e requerendo novas."
    ),
    "memoriais": "Memoriais para o julgamento em pauta, resumindo o caso a partir das fontes.",
    "habilitacao_credito": "Petição informando o óbito e requerendo a habilitação dos sucessores "
    "/ espólio (CPC, arts. 110, 313, I e § 2º, e 687 a 692).",
    "impugnacao_embargos_execucao": "Impugnação aos embargos à execução (CPC, art. 920, I).",
    "replica": "Réplica à contestação (CPC, arts. 350 e 351).",
    "excecao_pre_executividade": "Resposta à exceção de pré-executividade.",
    "resposta_agravo_interno": "Resposta ao agravo interno (CPC, art. 1.021, § 2º).",
}

SISTEMA = (
    "Você é advogado(a) do escritório Fonseca Brasil Serrão Advogados, credenciado do BANCO DA "
    "AMAZÔNIA S.A. (Basa), e redige um RASCUNHO de peça para revisão obrigatória por advogado "
    "inscrito na OAB antes do protocolo. Regras inegociáveis: (1) use somente fatos que estejam "
    "nas FONTES abaixo; o que faltar vira lacuna «[REVISOR: …]» (nome do juízo, valores, datas, "
    "subscritor, OAB); (2) não cite súmula, tema, acórdão ou número de precedente; cite apenas "
    "artigos de lei (CPC, Código Civil, Lei 10.931/2004, DL 167/1967) quando pertinentes; "
    "(3) o Banco é o credor/exequente salvo se as fontes disserem o contrário; (4) português "
    "jurídico formal e objetivo, sem adjetivar a parte contrária; (5) estrutura: endereçamento, "
    "qualificação resumida («[REVISOR]» onde faltar), fatos, fundamentos, pedidos, fecho com "
    "local/data/assinatura em lacuna; (6) no fim, uma seção '---\\nNotas para o revisor' com o "
    "que você presumiu e o que precisa ser conferido nos autos. Responda só com a peça em "
    "Markdown."
)


def _fontes(db: Session, o: MotorOportunidade) -> dict[str, Any]:
    f = db.scalar(select(BasaFicha).where(BasaFicha.cnj == o.cnj))
    dj = db.scalar(select(DataJudProcesso).where(DataJudProcesso.cnj == o.cnj))
    ultimos: list[str] = []
    if dj:
        movs = [m for g in (dj.graus or {}).values() for m in g.get("movimentos") or []]
        movs.sort(key=lambda m: m[2], reverse=True)
        ultimos = [
            f"{m[2]} — {m[1]}" + (f" ({', '.join(m[3])})" if m[3] else "") for m in movs[:15]
        ]
    prev = db.scalars(
        select(BasaAtoCobrado)
        .where(BasaAtoCobrado.cnj == o.cnj)
        .order_by(BasaAtoCobrado.competencia.desc())
    ).first()
    intim = db.scalars(
        select(TriagemEmail)
        .where(TriagemEmail.caixa == "djen", TriagemEmail.texto_integral.is_not(None))
        .order_by(TriagemEmail.recebido_em.desc())
    ).all()
    texto = next((t for t in intim if o.cnj in (t.cnjs or [])), None)
    return {
        "processo": _fmt(o.cnj),
        "ficha": f.ficha if f else o.ficha,
        "adverso": f.adverso if f else None,
        "valor_causa": f.valor_causa if f else None,
        "classe": dj.classe if dj else None,
        "comarca": prev.comarca if prev else None,
        "acao": prev.acao if prev else None,
        "evento": ROTULO.get(o.evento, o.evento),
        "data_evento": o.data_evento.isoformat(),
        "movimento": o.movimento,
        "observacao_motor": o.observacao,
        "ultimos_movimentos": ultimos,
        "intimacao_djen": (
            {
                "data": texto.recebido_em.date().isoformat() if texto.recebido_em else None,
                "assunto": texto.assunto,
                "texto": (texto.texto_integral or "")[:20000],
            }
            if texto
            else None
        ),
    }


def _fmt(d: str) -> str:
    return f"{d[:7]}-{d[7:9]}.{d[9:13]}.{d[13]}.{d[14:16]}.{d[16:]}" if len(d) == 20 else d


def gerar(
    db: Session,
    tenant_id: uuid.UUID,
    cnj: str,
    evento: str,
    data_evento: date,
    por: str,
    llm: Any | None = None,
) -> RascunhoPeca:
    o = db.scalar(
        select(MotorOportunidade).where(
            MotorOportunidade.cnj == cnj,
            MotorOportunidade.evento == evento,
            MotorOportunidade.data_evento == data_evento,
        )
    )
    if o is None:
        raise KeyError("oportunidade não está na última rodada do motor")
    if llm is None:
        from app.llm import get_llm_client

        llm = get_llm_client()
    fontes = _fontes(db, o)
    ato = carregar_tabela().atos.get(o.ato_id or "")
    instr = INSTRUCOES_POR_ATO.get(o.ato_id or "", "Peça adequada ao evento descrito.")
    linhas = [
        f"- {k}: {v}"
        for k, v in fontes.items()
        if k not in ("ultimos_movimentos", "intimacao_djen") and v
    ]
    if fontes["ultimos_movimentos"]:
        linhas.append(
            "- últimos movimentos no tribunal:\n  " + "\n  ".join(fontes["ultimos_movimentos"])
        )
    if fontes["intimacao_djen"]:
        i = fontes["intimacao_djen"]
        linhas.append(f"- intimação no DJEN ({i['data']}, {i['assunto']}):\n<<<\n{i['texto']}\n>>>")
    usuario = (
        f"PEÇA: {ATO.get(o.ato_id or '', o.ato_id)} — {ato.nome if ato else ''}\n"
        f"ORIENTAÇÃO: {instr}\n\nFONTES:\n" + "\n".join(linhas)
    )
    resp = llm.complete(
        [{"role": "system", "content": SISTEMA}, {"role": "user", "content": usuario}],
        purpose="rascunho_peca_basa",
        db=db,
        tenant_id=tenant_id,
        prompt_version=PROMPT_VERSAO,
    )
    corpo = re.sub(r"^```(?:markdown)?\s*|\s*```$", "", (resp.content or "").strip())
    aviso = (
        "> **RASCUNHO GERADO POR IA — revisão obrigatória por advogado antes do protocolo.** "
        f"Processo {fontes['processo']} · {fontes['evento']} em {o.data_evento:%d/%m/%Y}.\n\n"
    )
    r = RascunhoPeca(
        tenant_id=tenant_id,
        cnj=cnj,
        evento=evento,
        data_evento=data_evento,
        ato_id=o.ato_id,
        ficha=fontes["ficha"],
        conteudo=aviso + corpo,
        fontes={**fontes, "intimacao_djen": bool(fontes["intimacao_djen"])},
        modelo=getattr(resp, "provider_model", None) or getattr(resp, "model", None),
        status="gerado",
        criado_por=por[:80],
        criado_em=datetime.now(UTC),
    )
    db.add(r)
    db.commit()
    return r


def docx(r: RascunhoPeca) -> bytes:
    """Markdown simples → .docx (títulos, negrito, listas, citações)."""
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(12)

    def par(texto: str, estilo: str | None = None) -> None:
        p = doc.add_paragraph(style=estilo) if estilo else doc.add_paragraph()
        for i, pedaco in enumerate(re.split(r"\*\*", texto)):
            run = p.add_run(pedaco)
            run.bold = i % 2 == 1

    for linha in r.conteudo.splitlines():
        s = linha.rstrip()
        if not s:
            continue
        if s.startswith("#"):
            nivel = min(len(s) - len(s.lstrip("#")), 3)
            doc.add_heading(s.lstrip("#").strip(), level=nivel)
        elif s.startswith(("- ", "* ")):
            par(s[2:], "List Bullet")
        elif re.match(r"^\d+[.)] ", s):
            par(re.sub(r"^\d+[.)] ", "", s), "List Number")
        elif s.startswith(">"):
            par(s.lstrip("> ").strip(), "Intense Quote")
        elif s.strip() == "---":
            doc.add_paragraph("_" * 40)
        else:
            par(s)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def item(r: RascunhoPeca) -> dict[str, Any]:
    return {
        "id": str(r.id),
        "processo": _fmt(r.cnj),
        "evento": r.evento,
        "data_evento": r.data_evento.isoformat(),
        "ato": r.ato_id,
        "ficha": r.ficha,
        "conteudo": r.conteudo,
        "fontes": r.fontes,
        "modelo": r.modelo,
        "status": r.status,
        "criado_por": r.criado_por,
        "criado_em": r.criado_em.isoformat(),
    }
