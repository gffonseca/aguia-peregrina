"""Conciliação Benner × Legal One (Fase 3).

Entrada do Legal One: o export de processos (texto `CNJ|arquivado|pasta|basa`, uma linha por
processo — o formato que já usamos — ou planilha com as colunas Nº do processo / Situação /
Cliente). Quando a integração OData do Legal One estiver com credenciais no Railway, a mesma
função recebe os processos sincronizados.

Três listas, cada uma com a ação certa:
- ATIVA no Benner e AUSENTE no Legal One → cadastrar no Legal One (o escritório não controla);
- ATIVO no Legal One (cliente BASA) e AUSENTE no Benner → conferir com a CCONS (ficha não
  criada, ou processo que deveria ser encerrado no Legal One);
- ATIVO no Legal One e ENCERRADA no Benner → encerrar no Legal One.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models.faturamento import BasaFicha
from app.modules.m12_faturamento.tabela import normalizar

ATIVAS = {"em andamento", "em encerramento", "em reativacao"}


@dataclass
class ProcessoL1:
    cnj: str
    ativo: bool
    pasta: str
    basa: bool


def _dig(s: Any) -> str:
    return re.sub(r"\D", "", str(s or ""))


def ler_legalone(conteudo: bytes, nome: str = "") -> list[ProcessoL1]:
    out: list[ProcessoL1] = []
    if nome.lower().endswith((".xlsx", ".xlsm")) or conteudo[:2] == b"PK":
        ws = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True).worksheets[0]
        rows = list(ws.iter_rows(values_only=True))
        cab = [normalizar(str(c or "")) for c in rows[0]]
        i_n = next(i for i, c in enumerate(cab) if "processo" in c or c == "cnj")
        i_s = next((i for i, c in enumerate(cab) if "situac" in c or "status" in c), None)
        i_c = next((i for i, c in enumerate(cab) if "cliente" in c), None)
        i_p = next((i for i, c in enumerate(cab) if "pasta" in c), None)
        for r in rows[1:]:
            d = _dig(r[i_n])
            if len(d) != 20:
                continue
            sit = normalizar(str(r[i_s] or "")) if i_s is not None else ""
            cli = normalizar(str(r[i_c] or "")) if i_c is not None else "basa"
            out.append(
                ProcessoL1(
                    d,
                    not re.search(r"arquiv|encerr|baixad|inativ", sit),
                    str(r[i_p] or "") if i_p is not None else "",
                    bool(re.search(r"basa|banco da amazonia", cli)),
                )
            )
        return out
    for linha in conteudo.decode("utf-8", "replace").splitlines():
        p = linha.split("|")
        d = _dig(p[0]) if p else ""
        if len(d) != 20:
            continue
        arquivado = len(p) > 1 and p[1].strip().lower() in ("true", "1", "sim")
        basa = len(p) <= 3 or p[3].strip() in ("1", "true", "sim")
        out.append(ProcessoL1(d, not arquivado, p[2].strip() if len(p) > 2 else "", basa))
    return out


def conciliar(db: Session, l1: list[ProcessoL1]) -> tuple[bytes, dict[str, Any]]:
    fichas = list(db.scalars(select(BasaFicha)))
    ben_ativo = {f.cnj: f for f in fichas if normalizar(f.situacao) in ATIVAS and len(f.cnj) == 20}
    ben_todos = {f.cnj: f for f in fichas if len(f.cnj) == 20}
    l1_por = {p.cnj: p for p in l1}
    l1_basa_ativo = {p.cnj: p for p in l1 if p.ativo and p.basa}

    cadastrar_l1 = [f for c, f in ben_ativo.items() if c not in l1_por]
    conferir_ccons = [p for c, p in l1_basa_ativo.items() if c not in ben_todos]
    encerrar_l1 = [
        (p, ben_todos[c]) for c, p in l1_basa_ativo.items() if c in ben_todos and c not in ben_ativo
    ]

    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Resumo"
    resumo = {
        "benner_ativas": len(ben_ativo),
        "legalone_basa_ativos": len(l1_basa_ativo),
        "nos_dois": len(set(ben_ativo) & set(l1_por)),
        "cadastrar_no_legalone": len(cadastrar_l1),
        "conferir_com_ccons": len(conferir_ccons),
        "encerrar_no_legalone": len(encerrar_l1),
    }
    for k, v in resumo.items():
        ws.append([k.replace("_", " "), v])
    ws.column_dimensions["A"].width = 34

    def aba(titulo: str, cab: list[str], linhas: list[list[Any]]) -> None:
        a = wb.create_sheet(titulo)
        a.append(cab)
        for c in a[1]:
            c.font = Font(bold=True)
        for linha in linhas:
            a.append(linha)
        for col in "ABCD":
            a.column_dimensions[col].width = 28

    fmt = lambda d: f"{d[:7]}-{d[7:9]}.{d[9:13]}.{d[13]}.{d[14:16]}.{d[16:]}"  # noqa: E731
    aba(
        "Cadastrar no Legal One",
        ["Processo", "Ficha", "Situação Benner", "Adverso"],
        [[fmt(f.cnj), f.ficha, f.situacao, f.adverso] for f in cadastrar_l1],
    )
    aba(
        "Conferir com a CCONS",
        ["Processo", "Pasta Legal One"],
        [[fmt(p.cnj), p.pasta] for p in conferir_ccons],
    )
    aba(
        "Encerrar no Legal One",
        ["Processo", "Pasta Legal One", "Ficha", "Situação Benner"],
        [[fmt(p.cnj), p.pasta, f.ficha, f.situacao] for p, f in encerrar_l1],
    )
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), resumo
